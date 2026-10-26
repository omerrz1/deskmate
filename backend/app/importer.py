import csv
import io
import json
import re

from openpyxl import load_workbook

from . import countries, db, tabular, voices

FIELD_HINTS = {
    "name": ["name", "full name", "customer", "client", "fullname", "שם", "الاسم", "nombre", "nom"],
    "phone": ["phone", "mobile", "tel", "cell", "number", "טלפון", "هاتف", "teléfono", "telephone"],
    "email": ["email", "e-mail", "mail", "מייל", "بريد"],
    "company": ["company", "business", "organization", "organisation", "employer", "חברה", "شركة"],
    "language": ["language", "lang", "locale", "שפה", "لغة", "idioma", "langue"],
    "status": ["status"],
    "tags": ["stage", "tag", "segment", "group", "category", "tier"],
    "notes": ["note", "comment", "remark", "הערה", "ملاحظ"],
}
LANG_ALIASES = {
    "en": "English", "eng": "English", "ar": "Arabic", "he": "Hebrew", "heb": "Hebrew", "es": "Spanish", "fr": "French",
    "de": "German", "ru": "Russian", "pt": "Portuguese", "tr": "Turkish", "it": "Italian", "hi": "Hindi", "ur": "Urdu",
    "bn": "Bengali", "bangla": "Bengali", "zh": "Chinese", "mandarin": "Chinese", "cantonese": "Chinese", "ja": "Japanese",
    "ko": "Korean", "id": "Indonesian", "ms": "Malay", "tl": "Tagalog", "filipino": "Tagalog", "ta": "Tamil", "te": "Telugu",
    "th": "Thai", "vi": "Vietnamese", "fa": "Persian", "farsi": "Persian", "pa": "Punjabi",
}


def read_rows_info(filename, data: bytes):
    """-> (header, body, info). A CSV or TSV is read by tabular (encoding, separator, title lines, repeated headings); `info` says what it found and repaired."""
    kind = tabular.kind_of(filename, data)
    if kind in ("csv", "tsv"):
        try:
            t = tabular.read_table(filename, data, max_rows=50000, max_cols=200, overflow="cut")
        except tabular.EmptyTable:
            return [], [], {}
        body = [{h: c.strip() for h, c in zip(t.columns, r)} for r in t.rows]
        return t.columns, body, {"read_as": t.describe(), "notes": t.notes, "warnings": t.warnings}
    if kind == "xls":
        raise tabular.TableError(tabular.XLS_MESSAGE)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise tabular.TableError("I couldn't open that Excel file. It may be damaged or password-protected.")
    header, body = [], []
    for ws in wb.worksheets:
        iterator = (list(r[:200]) for r in ws.iter_rows(values_only=True))
        first = next((r for r in iterator if any(c not in (None, "") for c in r)), None)
        if first is None:
            continue
        sheet_header = [str(c).strip() if c is not None else f"col{i+1}" for i, c in enumerate(first)]
        for c in sheet_header:
            if c not in header:
                header.append(c)
        for row in iterator:
            if not any(c not in (None, "") for c in row):
                continue
            body.append({sheet_header[i]: tabular._cell(c).strip() for i, c in enumerate(row) if i < len(sheet_header)})
            if len(body) >= 50000:
                break
        if len(body) >= 50000:
            break
    return header, body, {"read_as": f"XLSX, {len(body):,} rows × {len(header)} columns", "notes": [], "warnings": []}


def read_rows(filename, data: bytes):
    header, body, _info = read_rows_info(filename, data)
    return header, body


def guess_mapping(header):
    mapping = {}
    for field, hints in FIELD_HINTS.items():
        for col in header:
            low = col.lower()
            if col in mapping.values():
                continue
            if any(h == low for h in hints) or any(h in low for h in hints):
                mapping[field] = col
                break
    return mapping


def norm_lang(v):
    v = (v or "").strip()
    if not v:
        return "English"
    return LANG_ALIASES.get(v.lower(), v[:1].upper() + v[1:])


def norm_phone(v):
    v = (v or "").strip()
    v = re.sub(r"\.0$", "", v)
    keep = re.sub(r"[^\d+]", "", v)
    if keep.startswith("00"):
        keep = "+" + keep[2:]
    elif keep.isdigit() and countries.international(keep):
        keep = "+" + keep                 # 923001234567 -> +923001234567, so the country and language are found
    return keep


def import_rows(dataset_name, filename, header, body, mapping, source_doc_id=None):
    if source_doc_id:
        existing = db.q1("SELECT id,row_count FROM datasets WHERE source_doc_id=?", (source_doc_id,))
        if existing:
            return {"dataset_id": existing["id"], "imported": existing["row_count"], "skipped": 0, "already_imported": True}
    ds_id = db.run("INSERT INTO datasets(name,filename,columns,row_count,created_at,source_doc_id) VALUES(?,?,?,?,?,?)",
                   (dataset_name, filename, json.dumps(header), 0, db.now(), source_doc_id))
    mapped_cols = set(mapping.values())
    rows, skipped = [], 0
    for source_row, r in enumerate(body, 1):
        g = lambda f: r.get(mapping[f], "") if f in mapping else ""
        name, phone = g("name"), norm_phone(g("phone"))
        if not name and not phone:
            skipped += 1
            continue
        extra = {k: v for k, v in r.items() if k not in mapped_cols and v != ""}
        if source_doc_id:
            extra["__source_row"] = source_row
        sheet_lang = norm_lang(g("language")) if g("language") else ""
        lang, src = voices.choose_language(phone, sheet_lang)
        info = countries.lookup(phone)
        rows.append((ds_id, name or "Unknown", phone, g("email"), g("company"), lang,
                     (g("status") or "new").lower(), g("tags"), g("notes"), json.dumps(extra, ensure_ascii=False), db.now(),
                     info["country"] if info else "", src))
    db.many("INSERT INTO contacts(dataset_id,name,phone,email,company,language,status,tags,notes,extra,created_at,country,lang_source) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    db.update("datasets", ds_id, row_count=len(rows))
    db.log("dataset", f"Imported {len(rows)} contacts from {filename}", dataset_id=ds_id, skipped=skipped)
    return {"dataset_id": ds_id, "imported": len(rows), "skipped": skipped}
