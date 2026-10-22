"""Knowledge bases: persona + context + key facts + searchable documents, attachable to any campaign."""
import io
import re
from pathlib import Path

from . import db

db.run("""CREATE TABLE IF NOT EXISTS knowledge_bases(
  id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, agent_name TEXT DEFAULT '', persona TEXT DEFAULT '', context TEXT DEFAULT '',
  key_facts TEXT DEFAULT '', next_steps TEXT DEFAULT '', contact TEXT DEFAULT '', created_at REAL)""")
db.run("""CREATE TABLE IF NOT EXISTS kb_docs(
  id INTEGER PRIMARY KEY AUTOINCREMENT, kb_id INTEGER, title TEXT, kind TEXT DEFAULT 'info', body TEXT, created_at REAL)""")
db.run("CREATE INDEX IF NOT EXISTS idx_kbdocs ON kb_docs(kb_id)")
db.run("CREATE TABLE IF NOT EXISTS kb_pron(id INTEGER PRIMARY KEY AUTOINCREMENT, kb_id INTEGER, term TEXT, language TEXT, respelling TEXT)")
for _stmt in ("ALTER TABLE kb_docs ADD COLUMN category TEXT DEFAULT 'Other'", "ALTER TABLE kb_docs ADD COLUMN source TEXT"):
    try:
        db.run(_stmt)
    except Exception:
        pass

_HAS_FTS = False
try:
    existed = bool(db.q1("SELECT name FROM sqlite_master WHERE type='table' AND name='kb_docs_search'"))
    db.run("CREATE VIRTUAL TABLE IF NOT EXISTS kb_docs_search USING fts5(title, body, content='kb_docs', content_rowid='id')")
    db.run("CREATE TRIGGER IF NOT EXISTS kb_docs_search_insert AFTER INSERT ON kb_docs BEGIN INSERT INTO kb_docs_search(rowid,title,body) VALUES(new.id,new.title,new.body); END")
    db.run("CREATE TRIGGER IF NOT EXISTS kb_docs_search_delete AFTER DELETE ON kb_docs BEGIN INSERT INTO kb_docs_search(kb_docs_search,rowid,title,body) VALUES('delete',old.id,old.title,old.body); END")
    db.run("CREATE TRIGGER IF NOT EXISTS kb_docs_search_update AFTER UPDATE OF title,body ON kb_docs BEGIN INSERT INTO kb_docs_search(kb_docs_search,rowid,title,body) VALUES('delete',old.id,old.title,old.body); INSERT INTO kb_docs_search(rowid,title,body) VALUES(new.id,new.title,new.body); END")
    if not existed:
        db.run("INSERT INTO kb_docs_search(kb_docs_search) VALUES('rebuild')")
    _HAS_FTS = True
except Exception:
    # FTS5 is optional on custom SQLite builds; the lexical scan below remains a safe fallback.
    _HAS_FTS = False

CATEGORIES = ["Products & services", "Pricing & payment", "Requirements", "How to get started", "FAQ", "Policies", "Objection handling",
              "Contact & locations", "Support", "Other"]


def guess_category(title, body="", kind="info"):
    """Keyword fallback used when the AI classifier is unavailable."""
    t = f"{title}".lower()
    b = body[:600].lower()
    rules = [("How to get started", ("how to apply", "apply", "get started", "sign up", "signup", "order", "book", "enrol", "onboarding")),
             ("Contact & locations", ("contact", "address", "location", "find us", "opening hours", "branch", "office")),
             ("Pricing & payment", ("price", "pricing", "fee", "cost", "payment", "invoice", "subscription", "plan", "discount", "finance")),
             ("Requirements", ("requirement", "eligib", "prerequisite", "what you need")),
             ("Policies", ("policy", "policies", "procedure", "complaint", "refund", "terms", "privacy", "cancel")),
             ("Support", ("support", "help", "warranty", "service desk", "training")),
             ("FAQ", ("faq", "question")), ("Objection handling", ("objection", "too expensive", "not sure"))]
    if kind in ("course", "product"):
        return "Products & services"
    for cat, kws in rules:
        if any(k in t for k in kws):
            return cat
    for cat, kws in rules:
        if any(k in b for k in kws):
            return cat
    return "Other"


DEFAULT_PERSONA = ("Super realistic, warm, quick and natural, like a helpful colleague on a real phone call. Short sentences (one or two at a time), authentic breathing pauses, human filler words ('mm-hmm', 'yeah', 'hmm...'), no jargon, no sales pressure. Never pushy; if the person isn't interested, thank them warmly and end the call.")


# one-off migration for knowledge bases created before categories existed
for _d in db.q("SELECT id,title,kind,body FROM kb_docs WHERE source IS NULL"):
    db.run("UPDATE kb_docs SET category=?, source='imported' WHERE id=?", (guess_category(_d["title"], _d["body"], _d["kind"]), _d["id"]))


def list_kbs():
    return db.q("""SELECT k.id, k.name, k.agent_name, (SELECT COUNT(*) FROM kb_docs d WHERE d.kb_id=k.id) docs,
                   (SELECT COUNT(*) FROM campaigns c WHERE c.kb_id=k.id) campaigns FROM knowledge_bases k ORDER BY k.id""")


def get_kb(kb_id):
    if not kb_id:
        return None
    return db.q1("SELECT * FROM knowledge_bases WHERE id=?", (kb_id,))


def kb_detail(kb_id):
    kb = get_kb(kb_id)
    if kb:
        kb["docs"] = db.q("SELECT id,title,kind,category,source,length(body) size,created_at FROM kb_docs WHERE kb_id=? ORDER BY category, title", (kb_id,))
        kb["categories"] = CATEGORIES
    return kb


def item_titles(kb_id, categories=None):
    """Titles of the products/services documents, so the caller knows what the company offers."""
    if categories and "Products & services" not in categories:
        return []
    return [d["title"] for d in db.q("SELECT title FROM kb_docs WHERE kb_id=? AND category='Products & services' ORDER BY title", (kb_id,))]


def _chunks(doc):
    # Bound every chunk, including pages without headings or paragraph breaks.
    body = doc["body"] or ""
    parts = [body[i:i + 1400] for i in range(0, len(body), 1200)]
    source = f" | source: {doc['source']}" if doc.get("source") else ""
    return [f"[{doc['title']}{source}]\n{p.strip()}" for p in parts]


def _tok(s):
    return set(re.findall(r"\w{3,}", s.lower()))


def search(kb_id, query, k=3, categories=None, category=None):
    q = _tok(query) - _STOP
    if not q or not kb_id:
        return []
    docs = []
    if _HAS_FTS:
        match = " OR ".join(f'"{term}"*' for term in sorted(q))
        where = ["kb_docs_search MATCH ?", "d.kb_id=?"]
        args = [match, kb_id]
        if categories:
            where.append("d.category IN (" + ",".join("?" for _ in categories) + ")")
            args.extend(categories)
        if category:
            where.append("d.category=?")
            args.append(category)
        try:
            docs = db.q(f"SELECT d.title,d.body,d.category,d.source FROM kb_docs_search JOIN kb_docs d ON d.id=kb_docs_search.rowid WHERE {' AND '.join(where)} ORDER BY bm25(kb_docs_search) LIMIT 120", args)
        except Exception:
            docs = []
    if not docs:
        docs = db.q("SELECT title,body,category,source FROM kb_docs WHERE kb_id=?", (kb_id,))
    scored = []
    for d in docs:
        d = dict(d)
        if categories and d["category"] not in categories:
            continue
        if category and d["category"] != category:
            continue
        title_hit = len(q & _tok(d["title"]))
        for c in _chunks(d):
            score = len(q & _tok(c)) + 3 * title_hit
            if score:
                scored.append((score, c))
    scored.sort(key=lambda x: -x[0])
    out, seen = [], set()
    for _, chunk in scored:
        normalized = chunk.casefold()
        if normalized in seen:
            continue
        seen.add(normalized)
        out.append(chunk[:1500])
        if len(out) >= max(1, min(12, int(k or 3))):
            break
    return out


def add_doc(kb_id, title, text, kind="info", category=None, source=None):
    return db.run("INSERT INTO kb_docs(kb_id,title,kind,body,category,source,created_at) VALUES(?,?,?,?,?,?,?)",
                  (kb_id, title.strip() or "Note", kind, text.strip(), category or guess_category(title, text, kind), source, db.now()))


def segment_text(title, text, max_chars=1800):
    """Split a long document into topic segments: by markdown/numbered headings first, else by paragraphs."""
    text = text.strip()
    if len(text) <= max_chars:
        return [(title, text)]
    parts = re.split(r"\n(?=#{1,3} |\d+[.)] [A-Z][^\n]{3,80}\n)", text)
    by_heading = len(parts) >= 2
    if not by_heading:
        parts = re.split(r"\n{2,}", text)
    segs, cur = [], ""
    for p in parts:
        p = p.strip()
        if not p:
            continue
        # headings stay separate segments (merge only tiny ones); plain paragraphs are packed up to max_chars
        if cur and ((not by_heading and len(cur) + len(p) > max_chars) or (by_heading and len(cur) > 40)):
            segs.append(cur)
            cur = ""
        cur += p + "\n\n"
        if by_heading and len(cur) > max_chars:
            segs.append(cur)
            cur = ""
    if cur.strip():
        segs.append(cur)
    out = []
    for i, sgm in enumerate(segs, 1):
        first = re.sub(r"^#+\s*", "", sgm.strip().split("\n", 1)[0]).strip()
        label = first[:70] if 3 < len(first) <= 70 and not first.endswith(".") else f"part {i}"
        out.append((f"{title}: {label}", sgm.strip()))
    return out


async def classify(items):
    """items: [(title, body)] -> list of categories (AI when available, keywords otherwise)."""
    cats = [guess_category(t, b) for t, b in items]
    try:
        from . import llm
        if not llm.has_key() or not items:
            return cats
        payload = [{"i": i, "title": t, "text": b[:240]} for i, (t, b) in enumerate(items[:60])]
        res = await llm.complete_json(
            f"Sort knowledge-base snippets into exactly one category each. Categories: {CATEGORIES}. "
            "Reply as JSON {\"categories\": [..one category per snippet, same order..]}.", __import__("json").dumps(payload, ensure_ascii=False))
        got = res.get("categories") or []
        for i, c in enumerate(got[:len(cats)]):
            if c in CATEGORIES:
                cats[i] = c
    except Exception:
        pass
    return cats


async def ingest(kb_id, title, text, source=None):
    """Segment a document into topics, classify them, store them. Returns the number of segments."""
    segs = segment_text(title, text)
    cats = await classify(segs)
    for (t, b), c in zip(segs, cats):
        add_doc(kb_id, t, b, "info", c, source or title)
    return len(segs)


def extract_text(filename, data: bytes):
    name = (filename or "").lower()
    if name.endswith(".docx"):
        # Read Word's XML directly so DOCX support does not depend on a heavyweight
        # office suite or an optional package being installed in production.
        import zipfile
        import xml.etree.ElementTree as ET
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        parts = []
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 2500 or sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("Word document expands beyond the supported size")
            names = ["word/document.xml"]
            names.extend(sorted(n for n in archive.namelist() if re.fullmatch(r"word/(?:header|footer)\d+\.xml", n)))
            names.extend(n for n in ("word/footnotes.xml", "word/endnotes.xml") if n in archive.namelist())
            for part in names:
                if part not in archive.namelist():
                    continue
                info = archive.getinfo(part)
                if info.file_size > 15 * 1024 * 1024:
                    continue
                root = ET.fromstring(archive.read(part))
                for paragraph in root.iter(ns + "p"):
                    chunks = []
                    for node in paragraph.iter():
                        if node.tag == ns + "t" and node.text:
                            chunks.append(node.text)
                        elif node.tag == ns + "tab":
                            chunks.append("\t")
                        elif node.tag in (ns + "br", ns + "cr"):
                            chunks.append("\n")
                    line = "".join(chunks).strip()
                    if line:
                        parts.append(line)
        return "\n".join(parts)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages[:300])
    if name.endswith((".xlsx", ".xls")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                if any(c not in (None, "") for c in r):
                    out.append(" | ".join("" if c is None else str(c) for c in r))
        return "\n".join(out)
    return data.decode("utf-8", errors="replace")


# ---------- pronunciation guide: how names/abbreviations must be WRITTEN so the voice says them right ----------
def pronunciations(kb_id):
    return db.q("SELECT id,term,language,respelling FROM kb_pron WHERE kb_id=? ORDER BY language, term", (kb_id,)) if kb_id else []


def set_pronunciations(kb_id, rows):
    db.run("DELETE FROM kb_pron WHERE kb_id=?", (kb_id,))
    for r in rows:
        if str(r.get("term", "")).strip() and str(r.get("respelling", "")).strip():
            db.run("INSERT INTO kb_pron(kb_id,term,language,respelling) VALUES(?,?,?,?)",
                   (kb_id, r["term"].strip()[:60], (r.get("language") or "English").strip(), r["respelling"].strip()[:80]))


def pron_guide(kb_id, language):
    rows = [r for r in pronunciations(kb_id) if (r["language"] or "").lower() == (language or "").lower()]
    return "; ".join(f"{r['term']} -> write as \"{r['respelling']}\"" for r in rows)


_STOP = {"the", "and", "for", "what", "does", "have", "your", "about", "that", "this", "there", "their", "with", "please", "can", "you", "how", "are"}


def _file_candidates(terms):
    """Index full file text once per workspace; triggers keep edits searchable."""
    try:
        exists = db.q1("SELECT name FROM sqlite_master WHERE name='company_files_search'")
        if not exists:
            db.run("CREATE VIRTUAL TABLE IF NOT EXISTS company_files_search USING fts5(title, body, content='documents', content_rowid='id')")
            db.run("CREATE TRIGGER IF NOT EXISTS company_files_insert AFTER INSERT ON documents BEGIN INSERT INTO company_files_search(rowid,title,body) VALUES(new.id,new.title,new.body); END")
            db.run("CREATE TRIGGER IF NOT EXISTS company_files_delete AFTER DELETE ON documents BEGIN INSERT INTO company_files_search(company_files_search,rowid,title,body) VALUES('delete',old.id,old.title,old.body); END")
            db.run("CREATE TRIGGER IF NOT EXISTS company_files_update AFTER UPDATE OF title,body ON documents BEGIN INSERT INTO company_files_search(company_files_search,rowid,title,body) VALUES('delete',old.id,old.title,old.body); INSERT INTO company_files_search(rowid,title,body) VALUES(new.id,new.title,new.body); END")
            db.run("INSERT INTO company_files_search(company_files_search) VALUES('rebuild')")
        match = " OR ".join(f'"{term}"*' for term in sorted(terms))
        rows = db.q("SELECT d.* FROM company_files_search JOIN documents d ON d.id=company_files_search.rowid WHERE company_files_search MATCH ? ORDER BY bm25(company_files_search)", (match,))
        # Older binary uploads may predate extracted body storage.
        rows.extend(db.q("SELECT * FROM documents WHERE COALESCE(body,'')='' AND COALESCE(file_path,'')<>''"))
        return rows
    except Exception:
        return db.q("SELECT * FROM documents ORDER BY updated_at DESC")


def search_workspace(query, agent=None, k=8):
    """Search complete company/employee file contents, retaining source locations."""
    from . import team, files
    terms = _tok(query) - _STOP
    if not terms:
        return []
    scored = []
    for doc in _file_candidates(terms):
        if not team.can_access_doc(agent, doc):
            continue
        body = doc.get("body") or ""
        if not body and doc.get("file_path"):
            try:
                body = extract_text(doc["title"], (db.data_dir() / "files" / Path(doc["file_path"]).name).read_bytes())
            except Exception:
                continue
        source = f"file {doc['id']}, drive {doc['agent_id']}/{files.path_of(doc.get('folder_id'))}"
        if not terms.intersection(_tok(doc["title"] + " " + body)):
            continue
        for chunk in _chunks({"title": doc["title"], "body": body, "source": source}):
            score = len(terms & _tok(chunk))
            if score:
                scored.append((score, chunk))
    scored.sort(key=lambda item: -item[0])
    return [chunk for _, chunk in scored[:k]]


async def lookup_company(query, call_id=None, kb_id=None):
    """One retrieval path for phone transports: internal evidence, then website."""
    from . import voices, webdemo
    import asyncio
    call = db.q1("SELECT * FROM calls WHERE id=?", (call_id,)) if call_id else None
    camp = db.q1("SELECT * FROM campaigns WHERE id=?", (call.get("campaign_id"),)) if call else None
    kb_id = (camp or {}).get("kb_id") or kb_id
    agent_id = (camp or {}).get("agent_id") or (call or {}).get("report_agent_id")
    agent = db.q1("SELECT * FROM agents WHERE id=?", (agent_id,)) if agent_id else None
    internal = search_workspace(query, agent, k=6)
    hits = []
    # Admin answers must be available regardless of the campaign's attached KB.
    for kb in db.q("SELECT DISTINCT kb_id FROM kb_docs WHERE source='answered by the admin'"):
        hits.extend(search(kb["kb_id"], query, k=3))
    hits.extend(internal)
    if kb_id:
        hits.extend(search(kb_id, query, k=4))
    website_id = voices.get_setting("company_website_kb_id", "")
    company_ws = db.company()  # the website knowledge is kept once, in the company's own workspace, and every workspace of the company searches it
    if website_id and (str(website_id) != str(kb_id) or company_ws != db.current()):
        with db.use(company_ws):
            hits.extend(search(website_id, query, k=6))
    terms = _tok(query) - _STOP
    coverage = max((len(terms & _tok(hit)) for hit in hits), default=0)
    if not hits or coverage < min(2, len(terms)):
        url = voices.get_setting("company_website", "").strip()
        if url:
            # Calls must not wait for a full crawl. Inspect relevant known pages concurrently.
            with db.use(company_ws):
                sources = db.q("SELECT DISTINCT source FROM kb_docs WHERE kb_id=?", (website_id,)) if website_id else []
            terms = _tok(query) - _STOP
            urls = sorted([r["source"] for r in sources if str(r["source"]).startswith("http")],
                          key=lambda u: -len(terms & _tok(u)))[:5]
            urls = list(dict.fromkeys([*urls, url]))
            async def read(u):
                try:
                    return await asyncio.wait_for(webdemo.read_one(u), 8)
                except Exception:
                    return None
            pages = await asyncio.gather(*(read(u) for u in urls))
            if not sources:
                from urllib.parse import urlparse
                base = webdemo._base_domain(urlparse(url if "://" in url else "https://" + url).hostname)
                links = {link for page in pages if page for link in page.get("links", []) if webdemo._in_scope(link, base)}
                ranked = sorted(links, key=lambda u: -len(terms & _tok(u)))[:5]
                pages.extend(await asyncio.gather(*(read(u) for u in ranked if u not in urls)))
            for page in pages:
                if page:
                    for chunk in _chunks({**page, "body": page["text"], "source": page["url"]}):
                        if terms & _tok(chunk):
                            hits.append(chunk)
    return list(dict.fromkeys(hits))[:16] or ["No verified answer found after searching company files, saved admin answers and website evidence. Use ask_admin now; do not guess. An admin answer is authoritative for this call."]
