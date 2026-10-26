"""Table files, read and written the way a careful person would: CSV, TSV and Excel all arrive as the same clean table.

A CSV is not one format. Excel in much of Europe writes semicolons; "Unicode text" is UTF-16 with tabs; a Windows export is Windows-1252, or 1255 (Hebrew) or 1256 (Arabic),
not UTF-8. There may be a BOM, a `sep=;` hint line, a title above the header, two columns with one name, a stray quote mark that swallows the rest of the file, or a phone
number Excel shortened to 4.47E+11. Read naively, that is mojibake or one giant column, and an employee would work on garbage without knowing.

read_table() finds out what the file is, reads it, and says what it did (`notes`) and what it could not put right (`warnings`), in plain words an employee can repeat.
to_csv() and to_xlsx() write files Excel opens correctly and that cannot run a formula in someone's spreadsheet. Nothing here touches the database."""
import codecs
import csv
import datetime
import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass, field

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from . import sheet_ops as ops

csv.field_size_limit(16 * 1024 * 1024)   # a stray quote can make the rest of a file one "field": read it, then repair it

TEXT_FORMATS = {".csv": "csv", ".tsv": "tsv", ".tab": "tsv"}
EXTENSIONS = (".csv", ".tsv", ".tab", ".xlsx", ".xlsm", ".xls")
DELIMITERS = {"comma": ",", "semicolon": ";", "tab": "\t", "pipe": "|", ",": ",", ";": ";", "\t": "\t", "\\t": "\t", "|": "|"}
DELIMITER_NAMES = {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}
ENCODING_ALIASES = {"excel": "cp1252", "windows": "cp1252", "ansi": "cp1252", "western": "cp1252", "latin1": "latin-1", "hebrew": "cp1255", "arabic": "cp1256", "cyrillic": "cp1251",
                    "russian": "cp1251", "turkish": "cp1254", "greek": "cp1253", "polish": "cp1250", "chinese": "gb18030", "japanese": "cp932", "korean": "cp949", "unicode": "utf-16",
                    "utf8": "utf-8", "mac": "mac_roman"}
ENCODING_LABELS = {"utf-8": "UTF-8", "utf-8-sig": "UTF-8", "utf-16": "UTF-16", "utf-32": "UTF-32", "cp1252": "Windows-1252 (Western European)", "cp1255": "Windows-1255 (Hebrew)",
                   "cp1256": "Windows-1256 (Arabic)", "cp1251": "Windows-1251 (Cyrillic)", "cp1254": "Windows-1254 (Turkish)", "cp1253": "Windows-1253 (Greek)",
                   "cp1250": "Windows-1250 (Central European)", "gb18030": "Chinese (GB18030)", "big5": "Chinese (Big5)", "cp932": "Japanese (Shift-JIS)", "cp949": "Korean",
                   "mac-roman": "Mac Roman"}
_BOMS = [(codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"), (codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16")]
_EXCEL_TEXT = re.compile(r'^="(.*)"$', re.S)                                        # ="07447781753": how an export keeps a leading zero
_FORMULA_START = ("=", "@", "\t", "\r")
_NUMBERISH = re.compile(r"^[+-]\s*\d[\d\s().,/\-]*$")                              # -5, +44 7447 781753, -1,200.50: not formulas


class TableError(ValueError):
    """A file that cannot be read; the message is for the admin. `status` is the HTTP status an upload should answer with."""

    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


class EmptyTable(TableError):
    """A file with nothing in it."""


XLS_MESSAGE = "That is an old Excel file (.xls), which I can't open. In Excel choose Save As → Excel Workbook (.xlsx) or CSV, and upload that."


@dataclass
class Table:
    columns: list
    rows: list
    format: str = "csv"                  # csv | tsv | xlsx
    encoding: str = ""                   # of a text file
    delimiter: str = ""
    sheet: str = ""                      # of a workbook: the one that was read
    sheets: list = field(default_factory=list)
    header_line: int = 1                 # the header was this many rows down (more than 1 = there was a title above it)
    notes: list = field(default_factory=list)       # what was repaired, in words
    warnings: list = field(default_factory=list)    # what could not be repaired and what to check

    def body(self):
        import json
        return json.dumps({"columns": self.columns, "rows": self.rows}, ensure_ascii=False)

    def describe(self):
        bits = [f"{len(self.rows):,} rows × {len(self.columns)} columns"]
        if self.format != "xlsx":
            bits.append(f"{DELIMITER_NAMES.get(self.delimiter, 'comma')}-separated")
            bits.append(ENCODING_LABELS.get(self.encoding, self.encoding))
        elif self.sheet:
            bits.append(f"sheet '{self.sheet}'")
        return f"{self.format.upper()}, " + ", ".join(bits)

    def meta(self, filename="", **extra):
        """What is remembered about where a sheet came from (documents.meta): its format, so it can go back out the same way, and what the reader said."""
        out = {"format": self.format, "filename": filename, "rows": len(self.rows), "columns": len(self.columns), "read_as": self.describe(), "notes": self.notes, "warnings": self.warnings}
        if self.format != "xlsx":
            out.update(encoding=self.encoding, delimiter=self.delimiter)
        if self.sheet:
            out.update(sheet=self.sheet, sheets=self.sheets)
        return {**out, **extra}


# ---------- what kind of file ----------
def is_table(name):
    return str(name or "").lower().endswith(EXTENSIONS)


def kind_of(name, data):
    """'csv' | 'tsv' | 'xlsx' | 'xls', from the bytes first (people rename files), then the name."""
    low = str(name or "").lower()
    if data[:4] == b"PK\x03\x04":
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    for ext, kind in TEXT_FORMATS.items():
        if low.endswith(ext):
            return kind
    return "csv"          # .csv, and a .xls/.txt that is really text (many systems export those)


def sha(body):
    return hashlib.sha1(body.encode("utf-8")).hexdigest()


# ---------- text encoding ----------
def codec_name(encoding):
    key = str(encoding or "").strip().lower().replace("_", "-")
    try:
        return codecs.lookup(ENCODING_ALIASES.get(key, key)).name
    except LookupError:
        raise TableError(f"I don't know the text encoding '{encoding}'. Try utf-8, windows-1252, windows-1255 (Hebrew), windows-1256 (Arabic) or utf-16.")


def _utf16_without_bom(data):
    sample = data[:2000]
    if len(sample) < 4 or b"\x00" not in sample:
        return ""
    n = len(sample) // 2 * 2
    evens, odds, half = sample[0:n:2].count(0), sample[1:n:2].count(0), n / 2
    if odds > half * 0.3 and evens < half * 0.02:
        return "utf-16-le"
    if evens > half * 0.3 and odds < half * 0.02:
        return "utf-16-be"
    return ""


# Code pages the detector may name, and what proves them. Windows-1252 is the default for anything in Latin script: Excel on Windows in the UK and Western Europe writes it,
# and "£", "é" and "ü" read as Polish or Japanese under other pages, so a different page needs evidence of its own.
_NON_LATIN = {"cp1251", "cp1253", "cp1255", "cp1256", "cp874", "koi8-r", "koi8-u", "iso8859-5", "iso8859-6", "iso8859-7", "iso8859-8", "gb18030", "gbk", "gb2312", "big5", "big5hkscs",
              "cp932", "shift-jis", "shift-jisx0213", "euc-jp", "iso2022-jp", "cp949", "euc-kr", "cp950"}
_ONLY_THESE = {"cp1250": "ąćęłńśźżĄĆĘŁŃŚŹŻěščřžůďťňĚŠČŘŽŮĎŤŇőűŐŰ", "iso8859-2": "ąćęłńśźżĄĆĘŁŃŚŹŻěščřžůďťňĚŠČŘŽŮĎŤŇőűŐŰ", "mac-latin2": "ąćęłńśźżĄĆĘŁŃŚŹŻěščřžůďťňĚŠČŘŽŮĎŤŇőűŐŰ",
               "cp1254": "ğĞşŞıİ", "iso8859-9": "ğĞşŞıİ"}


def _legacy(data):
    """Which old code page wrote these bytes (they are not UTF-8): the detector's pick, if there is real evidence for it, else Windows-1252."""
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(data[:400_000]).best()
        guess = codecs.lookup(best.encoding).name.replace("_", "-") if best and best.encoding else "cp1252"
    except Exception:
        return "cp1252"
    high = data.translate(None, bytes(range(128)))
    if guess in _NON_LATIN:
        # whole words made of high bytes (a Hebrew, Arabic, Russian or Chinese name) vs single accents in Latin text
        in_words = sum(len(m.group()) for m in re.finditer(rb"[\x80-\xff]{3,}", data))
        return guess if len(set(high)) >= 4 and in_words >= 0.5 * len(high) else "cp1252"
    if guess in _ONLY_THESE:
        try:
            text = data.decode(guess)
        except UnicodeDecodeError:
            return "cp1252"
        beyond = [c for c in text if ord(c) > 127]
        marks = [c for c in beyond if c in _ONLY_THESE[guess]]
        need = 3 if guess.startswith(("cp1250", "iso8859-2", "mac")) else 2
        return guess if len(set(marks)) >= need and len(marks) >= 0.25 * len(beyond) else "cp1252"
    return "cp1252"


def decode(data, encoding=None):
    """-> (text, encoding, note). The encoding is worked out from the bytes unless it is given."""
    if encoding:
        name = codec_name(encoding)
        return data.decode(name, errors="replace"), name, f"Read as {ENCODING_LABELS.get(name, name)}, as asked."
    for bom, name in _BOMS:
        if data.startswith(bom):
            return data.decode(name, errors="replace"), name, ""
    wide = _utf16_without_bom(data)
    if wide:
        return data.decode(wide, errors="replace"), "utf-16", "The file is UTF-16 (Excel's 'Unicode Text')."
    try:
        return data.decode("utf-8"), "utf-8", ""
    except UnicodeDecodeError:
        pass
    text = data.decode("utf-8", errors="replace")
    high = len(data.translate(None, bytes(range(128))))
    bad = text.count("�")
    if high and bad / high < 0.05:     # UTF-8 with a few damaged bytes
        return text, "utf-8", f"{bad} character(s) in the file are damaged and show as �."
    name = _legacy(data)
    try:
        return data.decode(name), name, f"The file is not UTF-8; it is {ENCODING_LABELS.get(name, name)}."
    except UnicodeDecodeError:
        return data.decode(name, errors="replace"), name, f"The file is not UTF-8; it looks like {ENCODING_LABELS.get(name, name)}, with a few bytes that do not fit."


# ---------- delimited text ----------
def _lines(text):
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _records(text, delimiter, quoting=csv.QUOTE_MINIMAL, limit=None):
    out = []
    for r in csv.reader(io.StringIO(text), delimiter=delimiter, quoting=quoting):
        out.append(r)
        if limit and len(out) >= limit:
            break
    return out


def _consistency(rows):
    widths = [len(r) for r in rows if any(c.strip() for c in r)]
    if not widths:
        return 0, 0
    width, n = Counter(widths).most_common(1)[0]
    return n / len(widths), width


def sniff_delimiter(text, prefer=""):
    """-> the character that separates the columns: the one that splits the first rows into the same number of pieces, at least two."""
    best, best_key = ",", (0, 0)
    sample = text[:300_000]
    for d in (",", ";", "\t", "|"):
        try:
            share, width = _consistency(_records(sample, d, limit=300))
        except csv.Error:
            continue
        if width < 2 or share < 0.5:
            continue
        key = (round(share, 1) + (0.05 if d == prefer else 0), width)
        if key > best_key:
            best, best_key = d, key
    return best


def _parse(text, delimiter, notes):
    rows = _records(text, delimiter)
    lines = [ln for ln in text.split("\n") if ln.strip()]
    if len(lines) > 20 and len(rows) < 0.8 * len(lines) and any(c.count("\n") >= 10 for r in rows for c in r):
        # an opening quote that never closes would make the rest of the file one cell: read line by line instead
        alt = [r for ln in lines for r in _records(ln, delimiter)]
        if len(alt) > 1.2 * len(rows) and _consistency(alt)[0] >= 0.9:
            notes.append(f"A stray quote mark would have swallowed {len(alt) - len(rows) + 1:,} rows into one cell, so I read the file line by line.")
            return alt
    return rows


# ---------- tidying a raw grid into a table ----------
def _used(row):
    n = len(row)
    while n and not str(row[n - 1]).strip():
        n -= 1
    return n


def _header_index(rows):
    """Where the header is. Title lines above it ('Leads report', 'Exported 2 Oct') fill fewer cells than the table under them."""
    used = [_used(r) for r in rows[:300]]
    filled = [u for u in used if u]
    if len(filled) < 3:
        return 0
    mode = Counter(filled).most_common(1)[0][0]
    if mode < 2:
        return 0
    for i, u in enumerate(used[:20]):
        if u >= max(2, 0.6 * mode):
            return i
    return 0


def _names(header, n, notes):
    """Clean, non-blank, unique column names."""
    out, seen, blanks, renamed = [], set(), [], []
    for i in range(n):
        raw = ops.clean_text(header[i] if i < len(header) else "")[:ops.HEADER_LEN]
        if not raw:
            raw = f"Column {i + 1}"
            blanks.append(raw)
        name, k = raw, 2
        while ops.fold(name) in seen:
            name = f"{raw} ({k})"[:ops.HEADER_LEN + 4]
            k += 1
        if name != raw and raw not in blanks:
            renamed.append(raw)
        seen.add(ops.fold(name))
        out.append(name)
    if blanks:
        notes.append(f"{len(blanks)} column(s) had no heading, so I named them {blanks[0]}" + (f" … {blanks[-1]}" if len(blanks) > 1 else "") + ".")
    if renamed:
        notes.append(f"Two columns were both called {renamed[0]!r}; the second is now {renamed[0] + ' (2)'!r}." if len(renamed) == 1 else
                     f"{len(renamed)} headings were repeated ({', '.join(repr(r) for r in renamed[:3])}…); the repeats are numbered (2), (3).")
    return out


def _tidy(raw, table, max_rows, max_cols, overflow, cell_len):
    """raw: list of rows (lists of text) -> table.columns / table.rows, with the repairs listed in table.notes / table.warnings."""
    notes, warnings = table.notes, table.warnings
    raw = [r for r in raw if _used(r)]
    if not raw:
        raise EmptyTable("That file has no rows in it.")
    h = _header_index(raw)
    if h:
        titles = [next((str(c).strip() for c in r if str(c).strip()), "") for r in raw[:h]]
        notes.append(f"Skipped {h} line(s) above the table ({'; '.join(repr(t[:40]) for t in titles[:2] if t)}).")
    table.header_line = h + 1
    header, body = raw[h], raw[h + 1:]
    width = max(_used(r) for r in raw[h:h + 1 + len(body)])
    wide = sum(1 for r in body if _used(r) > _used(header))
    if wide and _used(header) < width:
        notes.append(f"{wide:,} row(s) had more values than the header has headings; the extra columns are called Column {_used(header) + 1} and so on.")
    if width > max_cols:
        dropped = [i for i in range(max_cols, width) if any(i < len(r) and str(r[i]).strip() for r in body)]
        if dropped:
            warnings.append(f"The file has {width} columns and a sheet holds {max_cols}: columns {max_cols + 1} to {width} ({len(dropped)} with data) were left out. Split the file by column to keep them.")
        width = max_cols
    table.columns = _names(header, width, notes)
    if len(body) > max_rows:
        if overflow == "error":
            raise TableError(f"That file has {len(body):,} data rows and a sheet holds at most {max_rows:,}. Split it into smaller files.", 413)
        warnings.append(f"Only the first {max_rows:,} of {len(body):,} rows were read.")
        body = body[:max_rows]
    cut = unwrapped = 0
    rows = []
    for r in body:
        cells = []
        for c in r[:width]:
            c = c if isinstance(c, str) else str(c)
            if c[:1] == "=" and c[-1:] == '"':
                m = _EXCEL_TEXT.match(c)
                if m:
                    c, unwrapped = m.group(1), unwrapped + 1
            if cell_len and len(c) > cell_len:
                c, cut = c[:cell_len], cut + 1
            cells.append(c)
        if len(cells) < width:
            cells.extend([""] * (width - len(cells)))
        rows.append(cells)
    table.rows = rows
    if unwrapped:
        notes.append(f"Removed Excel's =\"…\" text wrapper from {unwrapped:,} cell(s) (it is how an export keeps a leading zero).")
    if cut:
        warnings.append(f"{cut:,} cell(s) were longer than {cell_len} characters and were cut.")
    _check_numbers(table)
    return table


def _check_numbers(table):
    """Phone numbers Excel turned into 4.47E+11 have lost their digits for good: say so, loudly, before anyone dials them."""
    for ci, col in enumerate(table.columns):
        n = sum(1 for r in table.rows if len(r[ci]) < 20 and ("E+" in r[ci] or "e+" in r[ci]) and ops.SCIENTIFIC.match(r[ci].strip()))
        if n:
            table.warnings.append(f"{n:,} value(s) in '{col}' look like 4.47E+11: Excel shortened those long numbers when the file was saved, and the real digits are gone. "
                                  "Ask for a fresh export (or format the column as Text before saving); do not dial these.")


# ---------- reading ----------
def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=" ", timespec="seconds")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.isoformat(timespec="seconds")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else "%.15g" % v
    return str(v)


def _read_workbook(data, wanted, table):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise TableError("I couldn't open that Excel file. It may be damaged or password-protected.")
    try:
        names = list(wb.sheetnames)
        order = [n for n in names if wanted and n.strip().lower() == str(wanted).strip().lower()] if wanted else []
        if wanted and not order:
            raise TableError(f"The workbook has no sheet called '{wanted}'. Its sheets are: {', '.join(names)}.")
        for name in order or names:
            grid = [[_cell(c) for c in row] for row in wb[name].iter_rows(values_only=True)]
            if any(_used(r) for r in grid):
                table.sheet, table.sheets = name, names
                if len(names) > 1:
                    table.notes.append(f"The workbook has {len(names)} sheets; I read '{name}'. The others are: {', '.join(n for n in names if n != name)}.")
                return grid
    finally:
        wb.close()
    raise EmptyTable("That workbook has no rows in it.")


def read_table(filename, data, *, encoding=None, delimiter=None, sheet=None, max_rows=ops.MAX_ROWS, max_cols=ops.MAX_COLS, overflow="error", cell_len=ops.CELL_LEN):
    """Any CSV, TSV or Excel file -> Table. overflow='error' refuses more than max_rows rows; 'cut' keeps the first max_rows and warns; cell_len=0 keeps long cells whole. Raises TableError."""
    kind = kind_of(filename, data)
    table = Table([], [], kind)
    if kind == "xls":
        raise TableError(XLS_MESSAGE)
    if kind == "xlsx":
        grid = _read_workbook(data, sheet, table)
        return _tidy(grid, table, max_rows, max_cols, overflow, cell_len)
    stripped = data.lstrip()[:200].lower()
    if stripped.startswith((b"<!doctype", b"<html", b"<table", b"<?xml")):
        raise TableError("That file is a web page or XML saved with a spreadsheet name, not a table. Open it in Excel and Save As → CSV or Excel Workbook.")
    text, table.encoding, note = decode(data, encoding)
    if note:
        table.notes.append(note)
    text = _lines(text.replace("\x00", "").replace("﻿", ""))
    if sum(1 for ch in text[:20_000] if ch < " " and ch not in "\n\t") > 40:
        raise TableError("That doesn't look like a CSV or text file; it has binary data in it.")
    prefer = "\t" if kind == "tsv" else ""
    hint = re.match(r'^"?sep=(.)"?[ \t]*\n', text)
    if hint:
        text = text[hint.end():]
        prefer = hint.group(1)
        table.notes.append(f"Dropped Excel's 'sep={hint.group(1)}' first line.")
    if delimiter:
        table.delimiter = DELIMITERS.get(str(delimiter).strip().lower())
        if not table.delimiter:
            raise TableError(f"I don't know the separator '{delimiter}'. Use comma, semicolon, tab or pipe.")
    else:
        table.delimiter = sniff_delimiter(text, prefer)
        if table.delimiter != ("\t" if kind == "tsv" else ","):
            table.notes.append(f"The columns are separated by a {DELIMITER_NAMES[table.delimiter]}, not a {'tab' if kind == 'tsv' else 'comma'}.")
    return _tidy(_parse(text, table.delimiter, table.notes), table, max_rows, max_cols, overflow, cell_len)


# ---------- writing ----------
def safe_cell(v):
    """Text that cannot run as a formula when the file is opened in Excel (a customer's name can start with '='). Numbers and phone numbers (-5, +447…) are left alone."""
    v = "" if v is None else str(v)
    if not v:
        return v
    if v.startswith(_FORMULA_START) or (v[0] in "+-" and not _NUMBERISH.match(v)):
        return "'" + v
    return v


def csv_stream(columns, rows, delimiter=",", bom=True, safe=True, batch=500):
    """CSV as bytes, piece by piece (a long list never sits in memory): UTF-8 with a BOM so Excel shows £, Hebrew and Arabic correctly, CRLF line ends, quotes only where needed."""
    buf = io.StringIO(newline="")
    w = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    conv = safe_cell if safe else (lambda v: "" if v is None else str(v))
    if bom:
        yield codecs.BOM_UTF8
    w.writerow([conv(c) for c in columns])
    for n, r in enumerate(rows, 1):
        w.writerow([conv(c) for c in r])
        if n % batch == 0:
            yield buf.getvalue().encode("utf-8")
            buf.seek(0)
            buf.truncate()
    yield buf.getvalue().encode("utf-8")


def to_csv(columns, rows, delimiter=",", bom=True, safe=True):
    return b"".join(csv_stream(columns, rows, delimiter, bom, safe))


def xlsx_value(ws, v):
    """A value for a write-only worksheet. Text is cleaned of characters Excel refuses, and text that starts with '=' is kept as text: openpyxl would otherwise write it as a formula,
    and a note like '=== VIP ===' is an invalid one, which makes Excel say the whole file is corrupt."""
    from openpyxl.cell import WriteOnlyCell
    if not isinstance(v, str):
        return v
    v = ILLEGAL_CHARACTERS_RE.sub("", v)
    if v.startswith("="):
        cell = WriteOnlyCell(ws, value=v)
        cell.data_type = "s"
        return cell
    return v


def to_xlsx(columns, rows):
    """An Excel workbook with every cell as text, header row frozen."""
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.freeze_panes = "A2"
    for r in [columns, *rows]:
        ws.append([xlsx_value(ws, "" if v is None else str(v)) for v in r])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
