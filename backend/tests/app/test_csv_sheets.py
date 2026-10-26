"""CSV is a first-class sheet: it uploads, the employee works on it as well as on Excel, and the result downloads in one click as CSV or Excel."""
import io
import json
import unittest

from openpyxl import load_workbook

from app import db, files, main, sheet_tools, tabular, voices, workbench
from .test_drives import DriveBase
from .test_sheet_tools import SheetBase

SEMICOLON = "Name;Mobile;E-mail;Fee\r\nJosé Müller;07447 781753;JOSE@X.COM;£1,200\r\nZoë O'Brien;92 300 1234567;zoe@x.com;£950\r\n".encode("cp1252")


def attachments(meeting_id):
    out = []
    for r in db.q("SELECT attachments FROM meeting_msgs WHERE meeting_id=? AND role='agent' AND attachments<>'[]' ORDER BY id", (meeting_id,)):
        out += json.loads(r["attachments"])
    return out


class UploadAndDownloadTests(DriveBase):
    async def upload(self, name, raw, who=None):
        return await (who or self.owner).upload("/api/fs/upload", name, raw, {"owner": "shared"})

    async def sheet(self, did):
        st, d = await self.owner.get(f"/api/docs/{did}")
        self.assertEqual(st, 200)
        return d

    async def test_a_semicolon_windows_csv_uploads_as_a_proper_sheet_and_says_what_it_did(self):
        st, r = await self.upload("leads.csv", SEMICOLON)
        self.assertEqual((st, r["kind"]), (200, "sheet"), r)
        self.assertIn("semicolon", r["summary"])
        self.assertIn("Windows-1252", r["summary"])
        d = await self.sheet(r["id"])
        self.assertEqual(d["table"]["columns"], ["Name", "Mobile", "E-mail", "Fee"])
        self.assertEqual(d["table"]["rows"][0], ["José Müller", "07447 781753", "JOSE@X.COM", "£1,200"])
        self.assertEqual((d["source"]["format"], d["source"]["delimiter"], d["source"]["encoding"]), ("csv", ";", "cp1252"))
        self.assertNotIn("meta", d)

    async def test_every_table_format_uploads(self):
        for name, raw in (("a.csv", b"Name,Phone\nAnn,1\n"), ("a.tsv", b"Name\tPhone\nAnn\t1\n"), ("a.tab", b"Name\tPhone\nAnn\t1\n"), ("a.xlsx", tabular.to_xlsx(["Name", "Phone"], [["Ann", "1"]]))):
            with self.subTest(name):
                st, r = await self.upload(name, raw)
                self.assertEqual((st, r["kind"]), (200, "sheet"), r)
                self.assertEqual((await self.sheet(r["id"]))["table"]["rows"], [["Ann", "1"]])

    async def test_unreadable_uploads_are_explained_not_500s(self):
        for name, raw, status, word in (("old.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, 400, ".xlsx"), ("empty.csv", b"\n\n", 400, "no rows"),
                                         ("big.csv", ("n\n" + "1\n" * 50_010).encode(), 413, "50,000"), ("page.xls", b"<html><table></table></html>", 400, "web page"),
                                         ("x.exe", b"MZ", 400, "Upload")):
            with self.subTest(name):
                st, r = await self.upload(name, raw)
                self.assertEqual(st, status, r)
                self.assertIn(word, r["detail"])

    async def test_an_untouched_upload_downloads_byte_for_byte(self):
        _, r = await self.upload("leads.csv", SEMICOLON)
        st, raw = await self.owner.get(f"/api/docs/{r['id']}/download")
        self.assertEqual((st, raw), (200, SEMICOLON))
        self.assertIn('filename="leads.csv"', self.owner.last_headers["content-disposition"])

    async def test_an_edited_sheet_downloads_what_is_in_it_now_not_the_stale_upload(self):
        """Regression: the original upload was always handed back, so everything the employee had cleaned or called into was lost on download."""
        _, r = await self.upload("leads.csv", SEMICOLON)
        d = await self.sheet(r["id"])
        table = d["table"]
        table["rows"][0][1] = "+447447781753"
        self.assertEqual((await self.owner.put(f"/api/docs/{r['id']}/table", table))[0], 200)
        st, raw = await self.owner.get(f"/api/docs/{r['id']}/download")
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "a BOM, so Excel shows £ and accents")
        text = raw.decode("utf-8-sig")
        self.assertIn("José Müller;+447447781753;JOSE@X.COM;£1,200", text)   # the semicolon the file came with is kept
        self.assertNotEqual(raw, SEMICOLON)
        self.assertIn('filename="leads.csv"', self.owner.last_headers["content-disposition"])     # never leads.csv.xlsx

    async def test_the_same_sheet_downloads_as_excel_or_csv_on_request(self):
        _, r = await self.upload("leads.csv", SEMICOLON)
        st, raw = await self.owner.get(f"/api/docs/{r['id']}/download", query="format=xlsx")
        self.assertEqual(st, 200)
        self.assertIn(".xlsx", self.owner.last_headers["content-disposition"])
        ws = load_workbook(io.BytesIO(raw)).active
        self.assertEqual([c.value for c in ws[2]], ["José Müller", "07447 781753", "JOSE@X.COM", "£1,200"])
        _, raw = await self.owner.get(f"/api/docs/{r['id']}/download", query="format=csv")
        self.assertEqual(raw.decode("utf-8-sig").splitlines()[0], "Name;Mobile;E-mail;Fee")
        _, raw = await self.owner.get(f"/api/docs/{r['id']}/download", query="format=tsv")
        self.assertEqual(raw.decode("utf-8-sig").splitlines()[1].split("\t")[0], "José Müller")
        st, err = await self.owner.get(f"/api/docs/{r['id']}/download", query="format=pdf")
        self.assertEqual(st, 400)

    async def test_an_excel_upload_still_downloads_as_excel(self):
        _, r = await self.upload("leads.xlsx", tabular.to_xlsx(["Name"], [["Ann"]]))
        await self.owner.put(f"/api/docs/{r['id']}/table", {"columns": ["Name"], "rows": [["Ann"], ["Bob"]]})
        _, raw = await self.owner.get(f"/api/docs/{r['id']}/download")
        self.assertEqual([c.value for (c,) in load_workbook(io.BytesIO(raw)).active.iter_rows(min_row=2)], ["Ann", "Bob"])
        self.assertIn('filename="leads.xlsx"', self.owner.last_headers["content-disposition"])

    async def test_hebrew_and_arabic_names_download_without_a_server_error(self):
        _, r = await self.upload("לקוחות.csv", "שם,טלפון\nדוד,0501234567\n".encode("utf-8"))
        await self.owner.put(f"/api/docs/{r['id']}/table", {"columns": ["שם"], "rows": [["דוד"]]})
        st, raw = await self.owner.get(f"/api/docs/{r['id']}/download")
        self.assertEqual(st, 200)
        self.assertIn("filename*=UTF-8''", self.owner.last_headers["content-disposition"])
        self.assertIn("דוד", raw.decode("utf-8-sig"))

    async def test_a_download_is_a_valid_file_whatever_is_in_the_cells(self):
        """Regression: a note such as '=== VIP ===' was written to Excel as an invalid formula, and a control character crashed the download; Excel called the file corrupt."""
        import zipfile
        st, made = await self.owner.post("/api/fs/file", {"owner": "shared", "title": "Notes", "kind": "sheet", "columns": ["Name", "Note"]})
        await self.owner.put(f"/api/docs/{made['id']}/table", {"columns": ["Name", "Note"], "rows": [["Ann", "=== VIP ==="], ["Bob", "=+"], ["Cy", "a\x0bb"], ["Di", "@SUM(1)"], ["דוד", "£1,200"]]})
        st, raw = await self.owner.get(f"/api/docs/{made['id']}/download")
        self.assertEqual(st, 200)
        archive = zipfile.ZipFile(io.BytesIO(raw))
        self.assertIsNone(archive.testzip())
        self.assertNotIn(b"<f>", archive.read("xl/worksheets/sheet1.xml"))            # no formulas anywhere
        values = [[c.value for c in r] for r in load_workbook(io.BytesIO(raw)).active.iter_rows(min_row=2)]
        self.assertEqual(values, [["Ann", "=== VIP ==="], ["Bob", "=+"], ["Cy", "ab"], ["Di", "@SUM(1)"], ["דוד", "£1,200"]])
        self.assertIn(".xlsx", self.owner.last_headers["content-disposition"])

    async def test_an_uploaded_csv_is_never_served_as_an_xlsx(self):
        """Regression: the file was the CSV as uploaded but the page saved it as <name>.csv.xlsx, which Excel refuses to open."""
        _, r = await self.upload("leads.csv", b"Name,Phone\nAnn,1\n")
        for query in ("", "format=csv"):
            st, raw = await self.owner.get(f"/api/docs/{r['id']}/download", query=query)
            self.assertNotEqual(raw[:2], b"PK")
            self.assertIn('filename="leads.csv"', self.owner.last_headers["content-disposition"])
            self.assertNotIn("xlsx", self.owner.last_headers["content-disposition"])
        st, raw = await self.owner.get(f"/api/docs/{r['id']}/download", query="format=xlsx")
        self.assertEqual(raw[:2], b"PK")
        self.assertIn('filename="leads.xlsx"', self.owner.last_headers["content-disposition"])

    async def test_the_campaign_excel_report_cannot_contain_formulas_either(self):
        import zipfile
        with db.use(self.ws):
            ds = db.run("INSERT INTO datasets(name,filename,columns,row_count,created_at) VALUES(?,?,?,?,?)", ("Leads", "leads.csv", "[]", 1, db.now()))
            db.run("INSERT INTO contacts(dataset_id,name,phone,language,status,created_at) VALUES(?,?,?,?,?,?)", (ds, "=== VIP ===", "+447700900123", "English", "new", db.now()))
            cid = db.run("INSERT INTO campaigns(name,objective,status,filters,max_attempts,created_at) VALUES(?,?,?,?,?,?)", ("C", "o", "completed", json.dumps({"dataset_id": ds}), 1, db.now()))
            response = main.campaign_export(cid)
            raw = b"".join([chunk async for chunk in response.body_iterator])
        self.assertNotIn(b"<f>", zipfile.ZipFile(io.BytesIO(raw)).read("xl/worksheets/sheet1.xml"))
        self.assertEqual(list(load_workbook(io.BytesIO(raw)).active.values)[1][0], "=== VIP ===")

    async def test_a_copy_of_a_csv_sheet_keeps_its_format(self):
        _, r = await self.upload("leads.csv", SEMICOLON)
        st, copy = await self.owner.post(f"/api/docs/{r['id']}/copy", {"owner": "shared", "folder_id": None})
        self.assertEqual(st, 200)
        self.assertEqual((await self.sheet(copy["id"]))["source"]["format"], "csv")

    async def test_the_customer_page_reads_a_csv_the_same_way(self):
        st, p = await self.owner.upload("/api/datasets/preview", "clientes.csv", "Nombre;Teléfono\nJosé;92 300 1234567\nZoë;07700 900111\n".encode("cp1252"))
        self.assertEqual(st, 200, p)
        self.assertEqual((p["columns"], p["mapping"]["phone"]), (["Nombre", "Teléfono"], "Teléfono"))
        self.assertIn("semicolon", p["read_as"])
        st, done = await self.owner.upload("/api/datasets/import", "clientes.csv", "Nombre;Teléfono\nJosé;92 300 1234567\nZoë;07700 900111\n".encode("cp1252"),
                                           {"mapping": json.dumps({"name": "Nombre", "phone": "Teléfono"})})
        self.assertEqual((st, done["imported"]), (200, 2), done)
        with db.use(self.ws):
            phones = {r["name"]: (r["phone"], r["country"]) for r in db.q("SELECT name,phone,country FROM contacts")}
        self.assertEqual(phones["José"], ("+923001234567", "Pakistan"))      # recognised as Pakistani, so it can be dialled and gets the right country

    async def test_campaign_results_download_as_csv_too(self):
        with db.use(self.ws):
            ds = db.run("INSERT INTO datasets(name,filename,columns,row_count,created_at) VALUES(?,?,?,?,?)", ("Leads", "leads.csv", "[]", 2, db.now()))
            db.run("INSERT INTO contacts(dataset_id,name,phone,email,language,status,created_at) VALUES(?,?,?,?,?,?,?)", (ds, "=EVIL()", "+447700900123", "a@x.com", "English", "new", db.now()))
            db.run("INSERT INTO contacts(dataset_id,name,phone,email,language,status,created_at) VALUES(?,?,?,?,?,?,?)", (ds, "דוד, כהן", "+972501234567", "b@x.com", "Hebrew", "new", db.now()))
            cid = db.run("INSERT INTO campaigns(name,objective,status,filters,max_attempts,created_at) VALUES(?,?,?,?,?,?)", ("C", "o", "completed", json.dumps({"dataset_id": ds}), 1, db.now()))
            response = main.campaign_export_csv(cid)
            raw = b"".join([chunk async for chunk in response.body_iterator])
        self.assertTrue(raw.startswith(b"\xef\xbb\xbfName,Phone,Email"))
        text = raw.decode("utf-8-sig")
        self.assertIn("'=EVIL()", text)                        # a customer's name cannot run a formula in the admin's Excel
        self.assertIn('"דוד, כהן",+972501234567', text)
        self.assertIn("not called", text)


class SheetToolCsvTests(SheetBase):
    def csv_file(self, title, raw, meta=None):
        """A CSV stored the way an export or an attachment is: a file document, not a sheet."""
        stored = db.data_dir() / "files"
        stored.mkdir(parents=True, exist_ok=True)
        (stored / "stored1.csv").write_bytes(raw)
        return db.run("INSERT INTO documents(agent_id,title,kind,body,folder_id,file_path,mime,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",
                      (self.aid, title, "file", "", files.folder_id(self.aid, "Downloads"), "stored1.csv", "text/csv", db.now(), db.now()))

    async def test_a_csv_file_opens_as_a_sheet_by_name_the_first_time_any_sheet_tool_asks(self):
        with db.use(self.ws):
            self.csv_file("export.csv", SEMICOLON)
            out = await self.tool("sheet_profile", title="export.csv")
            self.assertEqual(out["rows"], 2)
            self.assertEqual([c["column"] for c in out["columns"]], ["Name", "Mobile", "E-mail", "Fee"])
            self.assertIn("semicolon", out["source"]["read_as"])
            self.assertTrue(any("Windows-1252" in n for n in out["source"]["fixed_on_the_way_in"]))
            before = db.q1("SELECT COUNT(*) n FROM documents WHERE kind='sheet'")["n"]
            await self.tool("sheet_profile", title="export.csv")
            await self.tool("sheet_query", title="export.csv")
            self.assertEqual(db.q1("SELECT COUNT(*) n FROM documents WHERE kind='sheet'")["n"], before)     # opened once, not once per question

    async def test_by_id_too_and_a_real_error_for_something_that_is_not_a_table(self):
        with db.use(self.ws):
            did = self.csv_file("export.csv", SEMICOLON)
            self.assertEqual((await self.tool("sheet_profile", id=did))["rows"], 2)
            pdf = db.run("INSERT INTO documents(agent_id,title,kind,body,created_at,updated_at) VALUES(?,?,?,?,?,?)", (self.aid, "brief.pdf", "file", "", db.now(), db.now()))
            self.assertIn("error", await self.tool("sheet_profile", id=pdf))

    async def test_pasted_csv_text_becomes_a_sheet_without_retyping_it(self):
        with db.use(self.ws):
            out, _ = await files.run_tool(self.agent(), "create_spreadsheet", {"title": "Pasted", "csv": "Name;Phone\nAnn;07700 900111\nBob;07700 900222\n"})
            self.assertEqual((out["rows"], out["columns"]), (2, ["Name", "Phone"]))
            self.assertEqual(self.table("Pasted")[1][1], ["Bob", "07700 900222"])
            bad, _ = await files.run_tool(self.agent(), "create_spreadsheet", {"title": "Nothing"})
            self.assertIn("error", bad)

    async def test_a_misread_file_can_be_read_again_with_the_right_encoding(self):
        raw = "שם,טלפון\nדוד כהן,0501234567\n".encode("cp1255")
        with db.use(self.ws):
            did = self.csv_file("hebrew.csv", raw)
            out, _ = await files.run_tool(self.agent(), "create_spreadsheet", {"title": "Hebrew list", "from_file": did, "encoding": "windows-1255"})
            self.assertEqual(self.table("Hebrew list")[0], ["שם", "טלפון"])
            self.assertTrue(any("as asked" in n for n in out["what_i_fixed"]))
            again, _ = await files.run_tool(self.agent(), "create_spreadsheet", {"title": "Latin", "from_file": did, "encoding": "windows-1252"})
            self.assertEqual(again["rows"], 1)                                       # a second reading of the same file is its own sheet; nothing is lost

    async def test_wrong_separator_is_visible_in_the_profile_and_fixable(self):
        with db.use(self.ws):
            did = self.csv_file("one.csv", b"Name;Phone;City\nAnn;1;Leeds\nBob;2;York\nCy;3;Hull\n")
            out, _ = await files.run_tool(self.agent(), "create_spreadsheet", {"title": "Forced comma", "from_file": did, "delimiter": "comma"})
            profile = await self.tool("sheet_profile", title="Forced comma")
            self.assertTrue(any("wrong separator" in p and "semicolon" in p for p in profile["header_problems"]), profile["header_problems"])

    async def test_cleaning_a_csv_gives_a_csv_and_puts_a_download_button_in_the_chat(self):
        with db.use(self.ws):
            self.csv_file("export.csv", SEMICOLON)
            out = await self.tool("sheet_clean", title="export.csv", operations=[{"op": "trim"}, {"op": "phone"}, {"op": "email"}], country_code="44")
            saved = out["saved_as"]
            self.assertIn("download button", saved["delivered_to_chat"])
            doc = db.q1("SELECT * FROM documents WHERE id=?", (saved["id"],))
            self.assertEqual(files.sheet_meta(doc)["format"], "csv")
            self.assertEqual(files.sheet_meta(doc)["delimiter"], ";")
            sent = attachments(self.mid)
            self.assertEqual([(a["doc_id"], a["kind"], a["format"]) for a in sent], [(saved["id"], "sheet", "csv")])
            name, mime, raw = files.download_parts(doc)
            self.assertEqual((name, mime.split(";")[0]), ("export clean.csv", "text/csv"))
            self.assertIn("+447447781753", raw.decode("utf-8-sig"))

    async def test_a_dry_run_delivers_nothing_and_a_sheet_already_in_the_chat_is_not_posted_again(self):
        with db.use(self.ws):
            self.sheet("Leads", ["Name", "Phone"], [["Ann", "07447 781753"], ["Bob", "07700 900111"]])
            await self.tool("sheet_clean", title="Leads", operations=[{"op": "phone"}], country_code="44", dry_run=True)
            self.assertEqual(attachments(self.mid), [])
            first = await self.tool("sheet_clean", title="Leads", operations=[{"op": "phone"}], country_code="44", in_place=True)
            self.assertIn("now has a download button", first["saved_as"]["delivered_to_chat"])
            second = await self.tool("sheet_clean", title="Leads", operations=[{"op": "trim"}], in_place=True)
            self.assertIn("already in the chat", second["saved_as"]["delivered_to_chat"])
            self.assertEqual(len(attachments(self.mid)), 1)
            # an explicit "send it again" always posts
            did = first["saved_as"]["id"]
            await files.run_tool(self.agent(), "share_chat_file", {"id": did})
            self.assertEqual(len(attachments(self.mid)), 2)

    async def test_outside_a_live_chat_the_tool_gives_download_links_instead(self):
        with db.use(self.ws):
            self.sheet("Leads", ["Name", "Phone"], [["Ann", "07447 781753"]])
            out, _ = await sheet_tools.run_tool({**self.agent(), "_meeting_id": None}, "sheet_clean", {"title": "Leads", "operations": [{"op": "trim"}]})
            saved = out["saved_as"]
            self.assertNotIn("delivered_to_chat", saved)
            self.assertTrue(saved["download"]["csv"].endswith(f"/api/docs/{saved['id']}/download?format=csv"))

    async def test_merging_a_csv_with_a_workbook_is_excel_and_two_csvs_stay_csv(self):
        with db.use(self.ws):
            self.csv_file("a.csv", b"Name,Phone\nAnn,07447 781753\n")
            self.csv_file("b.csv", b"Name,Phone\nBob,07700 900111\n")
            out = await self.tool("sheet_merge", left="a.csv", right="b.csv", mode="stack")
            self.assertEqual(files.sheet_meta(db.q1("SELECT * FROM documents WHERE id=?", (out["saved_as"]["id"],)))["format"], "csv")
            self.sheet("book", ["Name", "Phone"], [["Cy", "07700 900333"]])
            db.run("UPDATE documents SET meta=? WHERE title='book'", (json.dumps({"format": "xlsx"}),))
            out = await self.tool("sheet_merge", left="a.csv", right="book", mode="stack")
            self.assertEqual(files.sheet_meta(db.q1("SELECT * FROM documents WHERE id=?", (out["saved_as"]["id"],)))["format"], "xlsx")

    async def test_exporting_a_sheet_as_csv_makes_an_excel_safe_file_and_hands_it_over(self):
        with db.use(self.ws):
            did = self.sheet("Leads", ["Name", "Fee"], [["=EVIL()", "£1,200"], ["דוד", "5"]])
            out, _ = await workbench.run_tool(self.agent(), "export_document", {"id": did, "format": "csv"})
            self.assertEqual(out["file"], "Leads.csv")
            self.assertIn("download button", out["delivered_to_chat"])
            exported = db.q1("SELECT * FROM documents WHERE id=?", (out["file_id"],))
            self.assertEqual((exported["kind"], exported["title"]), ("file", "Leads.csv"))
            raw = (db.data_dir() / "files" / exported["file_path"]).read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
            self.assertIn("'=EVIL()", raw.decode("utf-8-sig"))
            # ...and the file it made opens straight back up as a sheet
            self.assertEqual((await self.tool("sheet_profile", title="Leads.csv"))["rows"], 2)

    async def test_the_live_failure_numbers_with_their_own_country_code_survive_cleaning(self):
        """The agent was asked to clean a lead list and call it: 647 Pakistani numbers became +4492… and 718 calls failed. Both the raw and the already-damaged forms are put right."""
        rows = [["A", "92 300 1234567"], ["B", "923011234567"], ["C", "+44923021234567"], ["D", "07447 781753"], ["E", "00971 50 123 4567"], ["F", "4.47448E+11"]]
        with db.use(self.ws):
            voices.set_setting("default_calling_code", "44")
            self.sheet("Scholarship leads", ["Name", "Phone"], rows)
            profile = await self.tool("sheet_profile", title="Scholarship leads")
            problems = " ".join(profile["columns"][1]["problems"])
            self.assertIn("in front of another country's number", problems)
            self.assertIn("4.47E+11", problems)
            out = await self.tool("sheet_clean", title="Scholarship leads", operations=[{"op": "phone"}], country_code="44")
            cols, cleaned = self.table(out["saved_as"]["title"])
            self.assertEqual([r[1] for r in cleaned], ["+923001234567", "+923011234567", "+923021234567", "+447447781753", "+971501234567", "4.47448E+11"])
            self.assertEqual(out["rows_after"], 6)
            # a company with no calling code set: the damaged number is flagged with the instruction that repairs it, and the explicit code repairs it
            voices.set_setting("default_calling_code", "")
            self.sheet("Unset", ["Name", "Phone"], [["C", "+44923021234567"]])
            unset = await self.tool("sheet_profile", title="Unset")
            self.assertIn("country_code=44", " ".join(unset["columns"][1]["problems"]))
            fixed = await self.tool("sheet_clean", title="Unset", operations=[{"op": "phone"}], country_code="44")
            self.assertEqual(self.table(fixed["saved_as"]["title"])[1], [["C", "+923021234567"]])
            self.assertIn("Needs checking", cols)                                    # the shortened number is flagged, not guessed
            self.assertIn("scientific notation", cleaned[5][2])


if __name__ == "__main__":
    unittest.main()
