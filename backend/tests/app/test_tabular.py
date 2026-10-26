"""Reading and writing table files: a CSV from anywhere must arrive as the same clean table an Excel file does, and say what it repaired on the way."""
import io
import unittest

from openpyxl import Workbook, load_workbook

from app import countries, importer, knowledge, messaging, sheet_ops as ops, tabular

LEADS = (["Name", "Phone", "Amount"], [["José Müller", "07447 781753", "£1,200"], ["Zoë O'Brien", "07700 900111", "£950"], ["Renée Dupré", "07700 900333", "£80"]])


def xlsx(cols, rows, extra_sheets=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for name in extra_sheets:
        wb.create_sheet(name).append(["x"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def read(raw, name="leads.csv", **kw):
    return tabular.read_table(name, raw, **kw)


class EncodingTests(unittest.TestCase):
    def test_utf8_with_a_bom_has_no_bom_in_the_first_heading(self):
        t = read("﻿Name,Phone\nAnn,07700 900111\n".encode("utf-8"))
        self.assertEqual(t.columns, ["Name", "Phone"])
        self.assertEqual(t.rows, [["Ann", "07700 900111"]])

    def test_windows_1252_with_pound_signs_and_accents(self):
        text = "Name;Phone;Amount\n" + "\n".join(";".join(r) for r in LEADS[1]) + "\n"
        t = read(text.encode("cp1252"))
        self.assertEqual((t.encoding, t.delimiter), ("cp1252", ";"))
        self.assertEqual(t.rows[0], ["José Müller", "07447 781753", "£1,200"])
        self.assertTrue(any("Windows-1252" in n for n in t.notes))

    def test_a_few_pound_signs_alone_are_still_windows_1252_not_central_european(self):
        t = read("Name,Price\nA,£10\nB,£20\nC,£30\n".encode("cp1252"))
        self.assertEqual(t.rows[0][1], "£10")

    def test_many_accents_in_a_western_file_and_a_polish_file_are_each_right(self):
        west = "Name,City\nJosé Núñez,Zürich\nÉlodie Garçon,Besançon\nZoë Müller,Köln\nBjörn Åkesson,Malmö\nRenée Dupré,Orléans\n"
        self.assertEqual(read(west.encode("cp1252")).rows[1], ["Élodie Garçon", "Besançon"])
        polish = "Imię,Miasto\nŁukasz Żółć,Łódź\nZażółć Gęślą,Kraków\nJan Wiśniewski,Gdańsk\nMałgorzata Ćwik,Poznań\n"
        t = read(polish.encode("cp1250"))
        self.assertEqual((t.encoding, t.rows[0]), ("cp1250", ["Łukasz Żółć", "Łódź"]))

    def test_hebrew_windows_1255(self):
        text = "שם,טלפון,עיר\nדוד כהן,0501234567,תל אביב\nשרה לוי,0521234567,ירושלים\nיוסי מזרחי,0541234567,חיפה\n"
        t = read(text.encode("cp1255"))
        self.assertEqual(t.encoding, "cp1255")
        self.assertEqual(t.columns, ["שם", "טלפון", "עיר"])
        self.assertEqual(t.rows[1][2], "ירושלים")

    def test_arabic_windows_1256(self):
        text = "الاسم,الهاتف,المدينة\nمحمد علي,0501234567,القاهرة\nفاطمة حسن,0521234567,الرياض\nأحمد خالد,0541234567,دبي\n"
        t = read(text.encode("cp1256"))
        self.assertEqual(t.encoding, "cp1256")
        self.assertEqual(t.rows[0], ["محمد علي", "0501234567", "القاهرة"])

    def test_russian_turkish_chinese_and_japanese_are_found_too(self):
        for text, enc in (("Имя,Телефон,Город\nИван Петров,89161234567,Москва\nАнна Сидорова,89261234567,Санкт-Петербург\n", "cp1251"),
                          ("Ad,Telefon,Şehir\nAyşe Yılmaz,05321234567,İstanbul\nMehmet Öztürk,05421234567,Ankara\nGülşen Çelik,05521234567,İzmir\n", "cp1254"),
                          ("姓名,电话,城市\n王伟,13800138000,北京\n李娜,13900139000,上海\n张磊,13700137000,广州\n", "gbk"),
                          ("名前,電話,都市\n田中太郎,09012345678,東京\n鈴木花子,08012345678,大阪\n", "cp932")):
            with self.subTest(enc):
                t = read(text.encode(enc))
                self.assertEqual(t.columns[0], text.split(",")[0])
                self.assertEqual(t.rows[0][0], text.split("\n")[1].split(",")[0])

    def test_utf16_excel_unicode_text_with_tabs(self):
        t = read("Name\tPhone\nAnn\t07700 900111\nBob\t07700 900222\n".encode("utf-16"), "leads.txt.csv")
        self.assertEqual((t.encoding, t.delimiter), ("utf-16", "\t"))
        self.assertEqual(t.rows[1], ["Bob", "07700 900222"])

    def test_utf16_without_a_bom(self):
        t = read("Name,Phone\nAnn,07700 900111\nBob,07700 900222\n".encode("utf-16-le"))
        self.assertEqual(t.rows, [["Ann", "07700 900111"], ["Bob", "07700 900222"]])

    def test_mostly_utf8_with_a_damaged_byte_is_still_utf8(self):
        raw = ("Name,City\n" + "".join(f"José {i},Zürich\n" for i in range(200))).encode("utf-8") + b"Bad\xff,X\n"
        t = read(raw)
        self.assertEqual(t.encoding, "utf-8")
        self.assertEqual(t.rows[0], ["José 0", "Zürich"])
        self.assertTrue(any("damaged" in n for n in t.notes))

    def test_the_encoding_can_be_told_and_unknown_ones_are_refused(self):
        raw = "Name\nדוד\n".encode("cp1255")
        self.assertEqual(read(raw, encoding="hebrew").rows[0], ["דוד"])
        with self.assertRaises(tabular.TableError):
            read(raw, encoding="klingon")


class SeparatorAndShapeTests(unittest.TestCase):
    def test_separators_are_found(self):
        for sep in (",", ";", "\t", "|"):
            with self.subTest(sep):
                t = read(f"Name{sep}Phone{sep}City\nAnn{sep}07700 900111{sep}London\nBob{sep}07700 900222{sep}Leeds\n".encode())
                self.assertEqual(t.columns, ["Name", "Phone", "City"])

    def test_commas_inside_quotes_do_not_split_a_semicolon_file(self):
        t = read('Name;Address\nSmith;"1 High St, London"\nJones;"2 Low Rd, Leeds"\n'.encode())
        self.assertEqual((t.delimiter, t.rows[0]), (";", ["Smith", "1 High St, London"]))

    def test_excels_sep_hint_line_is_dropped_and_obeyed(self):
        t = read(b"sep=;\nName;Phone\nAnn;123456\n")
        self.assertEqual((t.columns, t.delimiter), (["Name", "Phone"], ";"))

    def test_a_tsv_extension_prefers_tabs(self):
        t = read(b"Name\tNote\nAnn\tliked it, a lot\n", "leads.tsv")
        self.assertEqual((t.format, t.delimiter, t.rows[0]), ("tsv", "\t", ["Ann", "liked it, a lot"]))

    def test_a_separator_can_be_forced(self):
        self.assertEqual(read(b"a;b\n1;2\n", delimiter="semicolon").columns, ["a", "b"])
        with self.assertRaises(tabular.TableError):
            read(b"a;b\n1;2\n", delimiter="slash")

    def test_title_lines_above_the_header_are_skipped(self):
        t = read(b"Leads report\nExported 2 Oct 2026\n\nName,Phone,Email,City\nAnn,1,a@x.com,Leeds\nBob,2,b@x.com,York\nCy,3,c@x.com,Hull\n")
        self.assertEqual(t.columns, ["Name", "Phone", "Email", "City"])
        self.assertEqual((len(t.rows), t.header_line), (3, 3))
        self.assertTrue(any("Skipped 2 line(s) above the table" in n for n in t.notes))

    def test_a_normal_first_row_is_never_taken_for_a_title(self):
        t = read(b"Name,Phone\nAnn,1\nBob,2\nCy,3\n")
        self.assertEqual((t.columns, len(t.rows), t.header_line), (["Name", "Phone"], 3, 1))

    def test_blank_and_repeated_headings_get_names(self):
        t = read(b"Name,,Phone,Name\nAnn,x,1,Annie\nBob,y,2,Bobby\n")
        self.assertEqual(t.columns, ["Name", "Column 2", "Phone", "Name (2)"])
        self.assertTrue(any("no heading" in n for n in t.notes))

    def test_a_repeated_heading_loses_no_data(self):
        t = read(b"Phone,Phone\n111,222\n")
        self.assertEqual(t.rows, [["111", "222"]])

    def test_rows_wider_than_the_header_get_columns_instead_of_losing_values(self):
        t = read(b"Name,Phone\nAnn,1,extra\nBob,2,more\nCy,3,x\n")
        self.assertEqual(t.columns, ["Name", "Phone", "Column 3"])
        self.assertEqual(t.rows[0], ["Ann", "1", "extra"])
        self.assertTrue(any("more values than the header" in n for n in t.notes))

    def test_empty_rows_and_trailing_empty_columns_go(self):
        t = read(b"Name,Phone,,\nAnn,1,,\n,,,\nBob,2,,\n")
        self.assertEqual((t.columns, t.rows), (["Name", "Phone"], [["Ann", "1"], ["Bob", "2"]]))

    def test_quoted_newlines_and_quotes_survive(self):
        t = read('Name,Note\nAnn,"line one\nline two"\nBob,"said ""hi"""\n'.encode())
        self.assertEqual(t.rows, [["Ann", "line one\nline two"], ["Bob", 'said "hi"']])

    def test_windows_and_old_mac_line_endings(self):
        for eol in ("\r\n", "\r"):
            with self.subTest(repr(eol)):
                self.assertEqual(read(f"Name,Phone{eol}Ann,1{eol}Bob,2{eol}".encode()).rows, [["Ann", "1"], ["Bob", "2"]])

    def test_excels_text_wrapper_around_a_leading_zero_is_removed(self):
        t = read(b'Name,Phone\nAnn,="07447781753"\nBob,="07700900111"\n')
        self.assertEqual(t.rows[0][1], "07447781753")
        self.assertTrue(any("text wrapper" in n for n in t.notes))

    def test_an_opening_quote_that_never_closes_does_not_swallow_the_file(self):
        body = "".join(f"Person {i},07700 {900000 + i},City {i}\n" for i in range(60))
        t = read(('Name,Phone,City\nBroken,"07700 1,Leeds\n' + body).encode())
        self.assertGreaterEqual(len(t.rows), 60)
        self.assertTrue(any("stray quote" in n for n in t.notes))

    def test_a_long_cell_is_cut_and_the_cut_is_reported(self):
        t = read(("Name,Note\nAnn," + "x" * 700 + "\n").encode())
        self.assertEqual(len(t.rows[0][1]), ops.CELL_LEN)
        self.assertTrue(any("were cut" in w for w in t.warnings))
        self.assertEqual(len(read(("Name,Note\nAnn," + "x" * 700 + "\n").encode(), cell_len=0).rows[0][1]), 700)

    def test_more_columns_than_a_sheet_holds_are_reported_not_hidden(self):
        t = read((",".join(f"C{i}" for i in range(45)) + "\n" + ",".join(str(i) for i in range(45)) + "\n").encode())
        self.assertEqual(len(t.columns), ops.MAX_COLS)
        self.assertTrue(any("columns 41 to 45" in w for w in t.warnings))

    def test_too_many_rows_is_refused_or_cut(self):
        raw = ("n\n" + "1\n" * 50_010).encode()
        with self.assertRaises(tabular.TableError) as ctx:
            read(raw)
        self.assertEqual(ctx.exception.status, 413)
        t = read(raw, overflow="cut")
        self.assertEqual(len(t.rows), 50_000)
        self.assertTrue(t.warnings)


class BadFileTests(unittest.TestCase):
    def test_nothing_in_the_file(self):
        for raw in (b"", b"\n\n,,\n"):
            with self.assertRaises(tabular.EmptyTable):
                read(raw)

    def test_an_old_xls_gets_a_clear_message(self):
        with self.assertRaises(tabular.TableError) as ctx:
            read(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "old.xls")
        self.assertIn(".xlsx", str(ctx.exception))

    def test_a_web_page_saved_as_xls_is_recognised(self):
        with self.assertRaises(tabular.TableError) as ctx:
            read(b"<html><table><tr><td>x</td></tr></table></html>", "report.xls")
        self.assertIn("web page", str(ctx.exception))

    def test_binary_is_refused(self):
        with self.assertRaises(tabular.TableError):
            read(bytes(range(1, 31)) * 300, "photo.csv")

    def test_a_tab_separated_file_named_xls_is_read_as_text(self):
        self.assertEqual(read(b"Name\tPhone\nAnn\t1\n", "export.xls").rows, [["Ann", "1"]])

    def test_a_csv_that_is_really_an_xlsx_is_read_as_one(self):
        t = read(xlsx(*LEADS), "renamed.csv")
        self.assertEqual((t.format, t.rows[0][0]), ("xlsx", "José Müller"))

    def test_a_damaged_workbook(self):
        with self.assertRaises(tabular.TableError):
            read(b"PK\x03\x04 not really a zip", "x.xlsx")


class WorkbookTests(unittest.TestCase):
    def test_numbers_dates_and_phone_floats_come_out_as_a_person_would_write_them(self):
        import datetime
        t = read(xlsx(["Name", "Phone", "Joined", "Qty", "Price", "Done"], [["Ann", 447447781753.0, datetime.datetime(2026, 10, 12), 5.0, 12.5, True]]), "x.xlsx")
        self.assertEqual(t.rows[0], ["Ann", "447447781753", "2026-10-12", "5", "12.5", "TRUE"])

    def test_other_sheets_are_named_and_a_sheet_can_be_chosen(self):
        raw = xlsx(*LEADS, extra_sheets=("Archive",))
        t = read(raw, "x.xlsx")
        self.assertEqual((t.sheet, t.sheets), ("Leads", ["Leads", "Archive"]))
        self.assertTrue(any("Archive" in n for n in t.notes))
        with self.assertRaises(tabular.TableError):
            read(raw, "x.xlsx", sheet="Nope")

    def test_csv_and_excel_give_the_same_table(self):
        """Mastery must not depend on the format: the same list in every dress reads identically."""
        cols, rows = LEADS
        dresses = {
            "xlsx": (xlsx(cols, rows), "a.xlsx"),
            "utf8": (tabular.to_csv(cols, rows), "a.csv"),
            "semicolon-1252": ("\r\n".join(";".join(r) for r in [cols, *rows]).encode("cp1252"), "a.csv"),
            "utf16-tab": ("\n".join("\t".join(r) for r in [cols, *rows]).encode("utf-16"), "a.txt.csv"),
            "pipe": ("\n".join("|".join(r) for r in [cols, *rows]).encode("utf-8"), "a.csv"),
        }
        for label, (raw, name) in dresses.items():
            with self.subTest(label):
                t = read(raw, name)
                self.assertEqual((t.columns, t.rows), (cols, rows))


class EvaluationWorldParityTests(unittest.TestCase):
    def test_the_messy_evaluation_lead_list_reads_identically_as_excel_and_as_every_kind_of_csv(self):
        """The live evaluation's 1,500-row lead list (stray spaces, mixed phone forms, switchboard numbers) must come out the same however it was saved."""
        from evals import data
        cols, rows = data.world()["leads"]
        rows = [[str(c) for c in r] for r in rows if any(str(c).strip() for c in r)]      # completely blank rows are dropped on the way in, for every format
        dresses = {"xlsx": (xlsx(cols, rows), "w.xlsx"), "utf8-comma": (tabular.to_csv(cols, rows), "w.csv"), "utf8-semicolon": (tabular.to_csv(cols, rows, ";"), "w.csv"),
                   "utf16-tab": (b"\xff\xfe" + "\r\n".join("\t".join(r) for r in [cols, *rows]).encode("utf-16-le"), "w.csv"), "tsv": (tabular.to_csv(cols, rows, "\t", bom=False), "w.tsv")}
        for label, (raw, name) in dresses.items():
            with self.subTest(label):
                t = read(raw, name)
                self.assertEqual((t.columns, len(t.rows)), (cols, len(rows)))
                self.assertEqual(t.rows, rows)
        # and the engine gives the same answer on all of them
        self.assertEqual(ops.profile(cols, read(dresses["utf8-semicolon"][0], "w.csv").rows, "44"), ops.profile(cols, read(dresses["xlsx"][0], "w.xlsx").rows, "44"))


class ScientificNotationTests(unittest.TestCase):
    def test_phones_excel_shortened_are_warned_about_loudly(self):
        t = read(b"Name,Phone\nAnn,4.47448E+11\nBob,7.44778E+11\nCy,07700 900111\n")
        self.assertEqual(len(t.warnings), 1)
        self.assertIn("2 value(s) in 'Phone'", t.warnings[0])
        self.assertIn("do not dial", t.warnings[0])

    def test_ordinary_scientific_numbers_are_left_alone(self):
        self.assertEqual(read(b"Name,Mass\nAnn,1.5E+3\n").warnings, [])

    def test_such_a_number_is_never_turned_into_a_phone_number(self):
        number, why = ops.norm_phone("4.47448E+11", "44")
        self.assertEqual(number, "")
        self.assertIn("scientific notation", why)
        cols, rows = ["Phone"], [["4.47448E+11"], ["07700 900111"]]
        issues = ops.profile(cols, rows, "44")["columns"][0]["problems"]
        self.assertTrue(any("4.47E+11" in p for p in issues))


class WritingTests(unittest.TestCase):
    def test_csv_opens_in_excel_with_the_right_characters(self):
        raw = tabular.to_csv(["שם", "Amount"], [["דוד, כהן", "£1,200"], ['say "hi"', "é"]])
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        text = raw.decode("utf-8-sig")
        self.assertEqual(text, 'שם,Amount\r\n"דוד, כהן","£1,200"\r\n"say ""hi""",é\r\n')

    def test_csv_round_trips(self):
        cols, rows = ["Name", "Note"], [["Ann", "line one\nline two"], ["Bob", 'a "b", c'], ["Zoë", ""]]
        for delimiter in (",", ";", "\t", "|"):
            with self.subTest(delimiter):
                t = read(tabular.to_csv(cols, rows, delimiter))
                self.assertEqual((t.columns, t.rows), (cols, rows))

    def test_a_customers_name_cannot_run_a_formula_in_someones_excel(self):
        raw = tabular.to_csv(["Name"], [["=HYPERLINK(\"http://evil\",\"x\")"], ["@SUM(1)"], ["+cmd|' /C calc'!A0"], ["-2+3"], ["\t=x"]]).decode("utf-8-sig")
        for line in raw.split("\r\n")[1:-1]:
            self.assertTrue(line.lstrip('"').startswith("'"), line)

    def test_numbers_and_phone_numbers_are_not_mangled_by_the_formula_guard(self):
        for v in ("+447447781753", "+44 (0) 7447-781753", "-5", "-1,200.50", "07447781753", "£5", "Anne-Marie"):
            self.assertEqual(tabular.safe_cell(v), v)

    def test_xlsx_keeps_equals_text_as_text(self):
        raw = tabular.to_xlsx(["A"], [["=1+1"], ["a\x0bb"], ["+44"]])
        cells = [c for (c,) in load_workbook(io.BytesIO(raw)).active.iter_rows(min_row=2)]
        self.assertEqual((cells[0].value, cells[0].data_type), ("=1+1", "s"))
        self.assertEqual(cells[1].value, "ab")          # a control character would have broken the whole download
        self.assertEqual(cells[2].value, "+44")

    def test_a_long_list_streams_in_pieces(self):
        pieces = list(tabular.csv_stream(["n"], ([str(i)] for i in range(1200)), batch=500))
        self.assertGreaterEqual(len(pieces), 4)
        self.assertEqual(len(read(b"".join(pieces)).rows), 1200)


class PhoneRegressionTests(unittest.TestCase):
    """The live failure: numbers like 92 300 1234567 were given the company's +44, becoming +4492…, which Twilio refused (error 21216)."""

    def test_a_number_with_its_own_country_code_keeps_it(self):
        for raw, want in (("92 300 1234567", "+923001234567"), ("923001234567", "+923001234567"), ("971501234567", "+971501234567"), ("972501234567", "+972501234567"),
                          ("15551234567", "+15551234567"), ("353871234567", "+353871234567"), ("0092 300 1234567", "+923001234567")):
            with self.subTest(raw):
                self.assertEqual(ops.norm_phone(raw, "44"), (want, ""))

    def test_local_numbers_still_get_the_company_code(self):
        for raw, want in (("07447 781753", "+447447781753"), ("7447781753", "+447447781753"), ("2079460958", "+442079460958"), ("447447781753", "+447447781753"), ("+44 7447 781753", "+447447781753")):
            with self.subTest(raw):
                self.assertEqual(ops.norm_phone(raw, "44"), (want, ""))

    def test_numbers_already_damaged_by_the_company_code_are_put_right(self):
        for raw, want in (("+44923001234567", "+923001234567"), ("+44 971 501234567", "+971501234567"), ("+4492 3001234567", "+923001234567")):
            with self.subTest(raw):
                self.assertEqual(ops.norm_phone(raw, "44"), (want, ""))

    def test_a_uk_number_with_a_stray_digit_is_flagged_not_repaired_into_another_country(self):
        number, why = ops.norm_phone("4474477817530", "44")
        self.assertEqual(number, "")
        self.assertIn("too long for a +44 number", why)

    def test_a_german_mobile_is_not_mistaken_for_a_damaged_number(self):
        self.assertEqual(ops.norm_phone("+4915123456789", "49"), ("+4915123456789", ""))

    def test_pandas_style_floats(self):
        self.assertEqual(ops.norm_phone("447447781753.0", "44"), ("+447447781753", ""))
        self.assertEqual(ops.key_phone("447447781753.0"), ops.key_phone("+44 7447 781753"))

    def test_the_profile_tells_the_agent_what_is_wrong_and_that_clean_fixes_it(self):
        cols, rows = ["Name", "Phone"], [["A", "+44923001234567"], ["B", "+44971501234567"], ["C", "07447 781753"], ["D", "+447700900111"]]
        report = ops.profile(cols, rows, "44")
        problems = " ".join(report["columns"][1]["problems"])
        self.assertIn("2 number(s) have +44 in front of another country's number", problems)
        self.assertIn("repair 2", " ".join(s["why"] for s in report["suggested_operations"]))
        cleaner = ops.Cleaner(cols, rows, "44").run([{"op": "phone"}])
        out_cols, out_rows, summary = cleaner.result()
        self.assertEqual([r[1] for r in out_rows], ["+923001234567", "+971501234567", "+447447781753", "+447700900111"])

    def test_the_dial_path_and_the_importer_agree(self):
        self.assertEqual(messaging.normalise("923001234567"), "+923001234567")
        self.assertEqual(importer.norm_phone("923001234567"), "+923001234567")
        self.assertEqual(importer.norm_phone("7447781753"), "7447781753")      # a local number is left for the dial step
        self.assertTrue(countries.international("971501234567"))
        self.assertFalse(countries.international("7447781753"))


class ProfileCatchesBadReadsTests(unittest.TestCase):
    def test_replacement_characters_and_mojibake_and_one_giant_column(self):
        report = ops.profile(["Name"], [["Jos�"], ["Ren�"], ["Zo�"]])
        self.assertTrue(any("wrong text encoding" in p for p in report["header_problems"]))
        report = ops.profile(["Name"], [["JosÃ©"], ["MÃ¼ller"]])
        self.assertTrue(any("Windows-1252" in p for p in report["header_problems"]))
        report = ops.profile(["Name;Phone;City"], [["Ann;1;Leeds"], ["Bob;2;York"], ["Cy;3;Hull"]])
        self.assertTrue(any("wrong separator" in p and "semicolon" in p for p in report["header_problems"]))


class KnowledgeAndImporterTests(unittest.TestCase):
    def test_a_csv_in_the_knowledge_base_is_readable_text(self):
        text = knowledge.extract_text("prices.csv", "Course;Price\nMBA;£9,500\nCyber;£4,200\n".encode("cp1252"))
        self.assertIn("MBA | £9,500", text)
        self.assertEqual(text.splitlines()[0], "Course | Price")

    def test_the_customer_import_reads_a_semicolon_windows_csv(self):
        raw = "Nombre;Teléfono;Idioma\nJosé;92 300 1234567;ur\nZoë;07700 900111;en\n".encode("cp1252")
        header, body, info = importer.read_rows_info("clientes.csv", raw)
        self.assertEqual(header, ["Nombre", "Teléfono", "Idioma"])
        self.assertEqual(body[0]["Nombre"], "José")
        self.assertIn("semicolon", info["read_as"])
        self.assertEqual(importer.guess_mapping(header)["phone"], "Teléfono")

    def test_repeated_headings_no_longer_overwrite_each_other_in_the_import(self):
        header, body = importer.read_rows("a.csv", b"Name,Phone,Phone\nAnn,111,222\n")
        self.assertEqual(header, ["Name", "Phone", "Phone (2)"])
        self.assertEqual(body[0]["Phone (2)"], "222")


if __name__ == "__main__":
    unittest.main()
