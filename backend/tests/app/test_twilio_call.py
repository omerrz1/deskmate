import asyncio
import base64
import hashlib
import io
import hmac
import json
import time
from types import SimpleNamespace
from unittest import mock
from xml.etree.ElementTree import fromstring
from openpyxl import load_workbook

import httpx
from app import chats, db, integrations, twilio_call, wa_call
from .test_people import Company


class TwilioFixture(Company):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.env = mock.patch.dict('os.environ', {'PUBLIC_BASE_URL': 'https://calls.example.com'})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.cfg = {'account_sid': 'AC' + 'a'*32, 'auth_token': 'test-secret', 'from_number': '+441234567890',
                    'whatsapp_from': '+441234567890', 'call_permission_content_sid': 'HX' + 'b'*32}
        with db.use(self.ws):
            db.run("INSERT INTO integrations(kind,name,config,enabled,allow_write) VALUES('twilio','Twilio',?,1,1)", (json.dumps(self.cfg),))
            self.chat = chats.get_or_create('whatsapp', 'twilio', '447700900123')
            db.update('chats', self.chat['id'], last_in_at=db.now())



class TwilioCalls(TwilioFixture):
    async def test_excel_import_reads_all_sheets_with_one_shared_lead_layout(self):
        from openpyxl import Workbook
        from app import importer
        wb = Workbook()
        first = wb.active
        first.append(["Name", "Phone"])
        first.append(["A", "+447700900101"])
        second = wb.create_sheet("Next batch")
        second.append(["Name", "Phone"])
        second.append(["B", "+447700900102"])
        raw = io.BytesIO()
        wb.save(raw)
        header, rows = importer.read_rows("leads.xlsx", raw.getvalue())
        self.assertEqual(header, ["Name", "Phone"])
        self.assertEqual([r["Name"] for r in rows], ["A", "B"])

    async def test_campaign_queue_handles_more_than_five_thousand_leads_in_batches(self):
        from app import callengine
        with db.use(self.ws), mock.patch.object(callengine, "_outbound_problem", return_value=""):
            dataset = db.run("INSERT INTO datasets(name,filename,columns,row_count,created_at) VALUES(?,?,?,?,?)", ("Large", "large.csv", "[]", 5505, db.now()))
            db.many("INSERT INTO contacts(dataset_id,name,phone,language,status,created_at) VALUES(?,?,?,?,?,?)",
                    ((dataset, f"Lead {i}", f"+447700{i:06d}", "English", "new", db.now()) for i in range(5505)))
            campaign = db.run("INSERT INTO campaigns(name,objective,filters,status,max_attempts,created_at) VALUES(?,?,?,?,?,?)",
                              ("Large batch", "Confirm interest", json.dumps({"dataset_id": dataset, "language": "English", "status": "new"}), "draft", 1, db.now()))
            self.assertEqual(callengine.queue_campaign(campaign), 5505)
            self.assertEqual(db.q1("SELECT COUNT(*) n FROM calls WHERE campaign_id=? AND status='queued'", (campaign,))["n"], 5505)
            self.assertEqual(callengine.queue_campaign(campaign), 0)

    async def test_new_workspace_gets_english_elevenlabs_voice_profile(self):
        from app import voices
        with db.use(self.ws):
            profiles = voices.profiles()
            english = next((p for p in profiles if p['language'] == 'English'), None)
            self.assertIsNotNone(english)
            self.assertTrue(english['voice_id'])

    async def test_unapproved_primary_profile_error_points_to_trust_hub_not_credentials(self):
        hint = integrations.twilio_error_hint(
            20003,
            "Primary compliance profile is not approved. Please refer to documentation and complete the KYC process in Trust Hub to gain access.",
            "voice")
        self.assertIn("Communications → Trust Hub → Profiles → Primary", hint)
        self.assertIn("Twilio Approved", hint)
        self.assertIn("not an Account SID or Auth Token error", hint)

    async def test_20003_checks_saved_credentials_before_advising_rotation(self):
        async def respond(req):
            if req.url.path.endswith('/Calls.json'):
                return httpx.Response(401, json={'code': 20003, 'message': 'Permission to access resource denied'}, request=req)
            return httpx.Response(200, json={'sid': self.cfg['account_sid'], 'status': 'active'}, request=req)
        with db.use(self.ws):
            async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
                with self.assertRaises(wa_call.CallError) as caught:
                    await twilio_call.request('Calls.json', {'To': '+447700900123'}, client, self.cfg)
        self.assertIn('authenticated the saved Account SID and Auth Token', str(caught.exception))
        self.assertIn('Do not replace the credentials', str(caught.exception))

    async def test_trial_recipient_errors_explain_verification_and_upgrade(self):
        hint = integrations.twilio_error_hint(572002, "No trial phone number assigned for this destination")
        self.assertIn("verify this recipient", hint)
        self.assertIn("upgrade the account", hint)
        call_hint = integrations.twilio_error_hint(21608, "trial recipient not verified", "voice")
        self.assertIn("Verified Caller IDs", call_hint)

    async def test_voice_rejection_keeps_provider_detail_even_when_code_is_zero(self):
        async def reject(req):
            return httpx.Response(400, json={"code": 0, "message": "Sender is not authorized for this route"}, request=req)
        with db.use(self.ws):
            async with httpx.AsyncClient(transport=httpx.MockTransport(reject)) as client:
                with self.assertRaisesRegex(wa_call.CallError, "Sender is not authorized for this route"):
                    await twilio_call.request("Calls.json", {"To": "+447700900123"}, client)

    async def test_voice_errors_map_trial_and_geo_blocks(self):
        trial = integrations.twilio_error_hint(21219, "To number not verified", "voice")
        geo = integrations.twilio_error_hint(21215, "Geo permissions", "voice")
        self.assertIn("verified recipient", trial)
        self.assertIn("Geo permissions", geo)

    async def test_permission_is_deduplicated_and_revocable(self):
        form = {'Body': 'VOICE_CALL_REQUEST', 'From': 'whatsapp:+447700900123', 'MessageSid': 'SM1', 'ButtonPayload': 'ACCEPTED'}
        with db.use(self.ws):
            self.assertTrue(twilio_call.permission_reply(form))
            first = chats._chat(self.chat['id'])['call_perm_until']
            twilio_call.permission_reply(form)
            self.assertEqual(first, chats._chat(self.chat['id'])['call_perm_until'])
            self.assertEqual(wa_call.perm_state(chats._chat(self.chat['id']))[0], 'granted')
            twilio_call.permission_reply({**form, 'MessageSid': 'SM2', 'ButtonPayload': 'REJECTED'})
            self.assertEqual(wa_call.perm_state(chats._chat(self.chat['id']))[0], 'denied')

    async def test_permission_request_uses_content_template(self):
        with db.use(self.ws), mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock, return_value={'sid': 'SM1'}) as send:
            await twilio_call.send_permission(self.chat)
            self.assertEqual(send.call_args.args[1]['ContentSid'], self.cfg['call_permission_content_sid'])
            self.assertEqual(send.call_args.args[1]['To'], 'whatsapp:+447700900123')

    async def test_call_routes_through_twilio_and_hangs_up_once(self):
        with db.use(self.ws):
            db.update('chats', self.chat['id'], call_perm='granted', call_perm_until=time.time()+3600)
            with mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock, return_value={'sid': 'CA'+'c'*32}) as send:
                result = await wa_call.start(self.chat['id'], None, 'Help the customer')
                self.assertIn('call_id', result, result)
                sess = wa_call.SESSIONS[(self.ws, result['call_id'])]
                self.assertIsInstance(sess, twilio_call.Session)
                payload = send.call_args.args[1]
                self.assertEqual(payload['From'], 'whatsapp:+441234567890')
                self.assertIn('/twilio/voice/media/', fromstring(payload['Twiml']).find('Connect/Stream').attrib['url'])
                self.assertEqual(payload['TimeLimit'], str(wa_call.MAX_MINUTES*60))
                self.assertEqual(payload['StatusCallbackEvent'], ['initiated', 'ringing', 'answered', 'completed'])
                await sess.end('no_answer')
                await sess.end('no_answer')
                self.assertEqual(send.await_count, 2)
                self.assertNotIn((self.ws, result['call_id']), wa_call.SESSIONS)

    async def test_phone_route_uses_twilio_voice_without_whatsapp_permission(self):
        with db.use(self.ws):
            from app import voices
            voices.set_setting('call_route', 'phone')
            with mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock, return_value={'sid': 'CA'+'d'*32}) as send:
                result = await wa_call.start(self.chat['id'], None, 'Confirm an appointment')
                self.assertIn('call_id', result, f"{result}: {db.q('SELECT message FROM activity ORDER BY id DESC LIMIT 3')}")
                self.assertEqual(send.call_args.args[1]['From'], self.cfg['from_number'])
                self.assertEqual(send.call_args.args[1]['To'], '+447700900123')
                self.assertNotIn('whatsapp:', send.call_args.args[1]['To'])
                self.assertEqual(db.q1('SELECT mode FROM calls WHERE id=?', (result['call_id'],))['mode'], 'phone')
                await wa_call.SESSIONS[(self.ws, result['call_id'])].end('no_answer')

    async def test_call_now_places_one_real_twilio_voice_call(self):
        from app import callengine
        with db.use(self.ws):
            contact_id = self.chat['contact_id']
            with mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock, return_value={'sid': 'CA'+'e'*32}) as send:
                result = await callengine.call_now(contact_id, 'Ask for an update')
                self.assertEqual(result['state'], 'initiated')
                payload = send.call_args.args[1]
                self.assertEqual(payload['From'], self.cfg['from_number'])
                self.assertEqual(payload['To'], '+447700900123')
                self.assertEqual(db.q1('SELECT status,mode,objective FROM calls WHERE id=?', (result['call_id'],)),
                                 {'status': 'ringing', 'mode': 'pstn', 'objective': 'Ask for an update'})
                await wa_call.SESSIONS[(self.ws, result['call_id'])].end('no_answer')

    async def test_campaign_dials_up_to_its_configured_concurrency_per_scheduler_tick(self):
        from app import callengine, agents
        with db.use(self.ws):
            campaign_id = db.run("INSERT INTO campaigns(name,objective,status,filters,max_attempts,concurrency,created_at) VALUES(?,?,?,?,?,?,?)",
                                 ('Batch', 'Confirm details', 'active', '{}', 1, 3, db.now()))
            for i in range(4):
                contact_id = db.run("INSERT INTO contacts(name,phone,language,status,created_at) VALUES(?,?,?,?,?)",
                                    (f'Lead {i}', f'+44770090012{i}', 'English', 'new', db.now()))
                db.run("INSERT INTO calls(contact_id,campaign_id,attempt,status,language,created_at,not_before,reason,mode) VALUES(?,?,?,?,?,?,?,?,?)",
                       (contact_id, campaign_id, 1, 'queued', 'English', db.now(), db.now(), 'initial', 'pstn'))

            class FakeSession:
                def __init__(self, ws, call_id, *args, **kwargs):
                    self.call_id = call_id
                async def place(self):
                    pass

            with mock.patch.object(twilio_call, 'voice_config', return_value=self.cfg), \
                 mock.patch.object(twilio_call, 'Session', FakeSession), \
                 mock.patch.object(agents, 'get', return_value=None):
                await callengine._dial_twilio_due()
            rows = db.q("SELECT status FROM calls WHERE campaign_id=? ORDER BY id", (campaign_id,))
            self.assertEqual([r['status'] for r in rows], ['ringing', 'ringing', 'ringing', 'queued'])

    async def test_campaign_excel_export_includes_latest_call_outcome(self):
        from app import main
        with db.use(self.ws):
            dataset_id = db.run("INSERT INTO datasets(name,filename,columns,row_count,created_at) VALUES(?,?,?,?,?)",
                                ('Leads', 'leads.xlsx', '[]', 1, db.now()))
            contact_id = db.run("INSERT INTO contacts(dataset_id,name,phone,email,company,language,status,created_at) VALUES(?,?,?,?,?,?,?,?)",
                                (dataset_id, 'A Lead', '+447700900123', 'lead@example.com', 'Example', 'English', 'contacted', db.now()))
            db.many("INSERT INTO contacts(dataset_id,name,phone,language,status,created_at) VALUES(?,?,?,?,?,?)",
                    ((dataset_id, f'Bulk {i}', f'+447701{i:06d}', 'English', 'new', db.now()) for i in range(10001)))
            campaign_id = db.run("INSERT INTO campaigns(name,objective,status,filters,max_attempts,created_at) VALUES(?,?,?,?,?,?)",
                                 ('Lead calls', 'Ask for update', 'completed', json.dumps({'dataset_id': dataset_id}), 1, db.now()))
            db.run("INSERT INTO calls(contact_id,campaign_id,attempt,status,outcome,summary,follow_up,report,created_at,ended_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                   (contact_id, campaign_id, 1, 'completed', 'callback_requested', 'Asked to call next week.', '2026-10-28 10:00',
                    json.dumps({"executive_summary": "Requested a callback next week.", "commitments": ["Team to call next week."], "objections_or_concerns": ["Timing uncertain."]}), db.now(), db.now()))
            response = main.campaign_export(campaign_id)
            raw = b''.join([chunk async for chunk in response.body_iterator])
            sheet = load_workbook(io.BytesIO(raw), read_only=True).active
            rows = list(sheet.values)
            self.assertEqual(rows[1][0:2], ('A Lead', '+447700900123'))
            self.assertEqual(rows[1][7], 'callback_requested')
            self.assertEqual(rows[1][9], 'Asked to call next week.')
            self.assertIn('Requested a callback next week.', rows[1][14])
            self.assertIn('Commitments', rows[1][14])
            self.assertEqual(len(rows), 10003)  # header plus every lead; exports are no longer capped at 10,000

    async def test_call_contact_accepts_a_phone_without_an_existing_contact_or_name(self):
        from app import admin_agent, callengine
        with db.use(self.ws), mock.patch.object(callengine, 'call_now', new_callable=mock.AsyncMock,
                                                return_value={'call_id': 44, 'state': 'ringing'}) as dial:
            result, link = await admin_agent.run_tool('call_contact',
                {'phone': '+447700900124', 'objective': 'Ask for an update'})
            self.assertEqual(result['state'], 'ringing')
            contact_id = dial.call_args.args[0]
            contact = db.q1('SELECT name,phone FROM contacts WHERE id=?', (contact_id,))
            self.assertEqual(contact['phone'], '+447700900124')
            self.assertTrue(contact['name'])
            self.assertEqual(link, '/calls?open=44')

    async def test_denied_permission_and_do_not_call_block_dial(self):
        with db.use(self.ws), mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock) as send:
            result = await wa_call.start(self.chat['id'], None, 'Help')
            self.assertIn('not allowed', result['error'])
            contact = chats.contact_for_phone(self.chat['ext_id'], '')
            db.update('contacts', contact, do_not_call=1)
            db.update('chats', self.chat['id'], contact_id=contact, call_perm='granted', call_perm_until=time.time()+3600)
            result = await wa_call.start(self.chat['id'], None, 'Help')
            self.assertIn('asked not', result['error'])
            send.assert_not_awaited()

    async def test_signature_binds_public_url_and_fields(self):
        path = f'/api/hooks/chat/{self.ws}/secret/twilio/voice'
        form = {'AccountSid': self.cfg['account_sid'], 'CallSid': 'CA123'}
        data = 'https://calls.example.com'+path+''.join(k+form[k] for k in sorted(form))
        sig = base64.b64encode(hmac.new(b'test-secret', data.encode(), hashlib.sha1).digest()).decode()
        req = SimpleNamespace(url=SimpleNamespace(path=path, query=''), headers={'x-twilio-signature': sig})
        self.assertTrue(twilio_call.signature_ok(req, form, self.cfg))
        self.assertFalse(twilio_call.signature_ok(req, {**form, 'CallSid': 'CA456'}, self.cfg))

    async def test_media_stream_accepts_twilio_websocket_trailing_slash_signature(self):
        path = f'/api/hooks/chat/{self.ws}/secret/twilio/voice/media/1'
        data = 'https://calls.example.com' + path + '/'
        sig = base64.b64encode(hmac.new(b'test-secret', data.encode(), hashlib.sha1).digest()).decode()
        socket = SimpleNamespace(url=SimpleNamespace(path=path, query=''), headers={'x-twilio-signature': sig})
        self.assertTrue(twilio_call.signature_ok(socket, {}, self.cfg, websocket=True))
        self.assertFalse(twilio_call.signature_ok(socket, {}, self.cfg))

    async def test_timeout_never_retries_call_creation(self):
        hits = []
        def timeout(req):
            hits.append(req)
            raise httpx.ReadTimeout('timed out', request=req)
        with db.use(self.ws):
            async with httpx.AsyncClient(transport=httpx.MockTransport(timeout)) as client:
                with self.assertRaisesRegex(wa_call.CallError, 'may have been accepted'):
                    await twilio_call.request('Calls.json', {'To': 'whatsapp:+447700900123'}, client)
        self.assertEqual(len(hits), 1)

    async def test_audio_encodes_and_interrupt_clears_playback(self):
        out = twilio_call.AudioOutput()
        out.push(b'\0\0' * 2400)
        kind, audio, generation = out.queue.get_nowait()
        self.assertEqual(kind, 'media')
        self.assertEqual(len(audio), 160)
        decoder = wa_call.av.CodecContext.create('pcm_mulaw', 'r')
        decoder.sample_rate = 8000
        decoder.layout = 'mono'
        self.assertEqual(decoder.decode(wa_call.av.Packet(audio))[0].samples, 160)
        out.pending.add('1')
        self.assertTrue(out.speaking)
        out.clear()
        self.assertFalse(out.pending)
        self.assertEqual(out.queue.get_nowait()[0], 'clear')
        self.assertFalse(out.speaking)

    async def test_stream_audio_and_playback_marks(self):
        with db.use(self.ws), mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock, return_value={'sid': 'CA'+'c'*32}):
            db.update('chats', self.chat['id'], call_perm='granted', call_perm_until=time.time()+3600)
            result = await wa_call.start(self.chat['id'], None, 'Help')
            sess = wa_call.SESSIONS[(self.ws, result['call_id'])]
            sess.engine = SimpleNamespace(send_audio=mock.AsyncMock(), close=mock.AsyncMock())
            await sess.receive_audio(base64.b64encode(b'\xff'*160).decode())
            self.assertGreater(len(sess.engine.send_audio.call_args.args[0]), 800)
            socket = SimpleNamespace(send_json=mock.AsyncMock(), close=mock.AsyncMock())
            sess.socket, sess.stream_sid = socket, 'MZ1'
            sess.out.push(b'\0\0'*2400)
            sender = asyncio.create_task(sess.send_output())
            await asyncio.sleep(0.001)
            self.assertEqual(socket.send_json.call_args_list[0].args[0]['event'], 'media')
            self.assertTrue(sess.out.pending)
            sess.out.clear()
            await asyncio.sleep(0.03)
            self.assertIn('clear', [c.args[0]['event'] for c in socket.send_json.call_args_list])
            sender.cancel()
            await asyncio.gather(sender, return_exceptions=True)
            await sess.end('no_answer')

    async def test_callback_requires_signature_and_same_account(self):
        from starlette.requests import Request
        from urllib.parse import urlencode
        from fastapi import HTTPException
        with db.use(self.ws):
            path = f'/api/hooks/chat/{self.ws}/{chats.secret()}/twilio/voice'
            form = {'AccountSid': self.cfg['account_sid'], 'CallSid': 'CA'+'c'*32}
            async def receive():
                return {'type': 'http.request', 'body': urlencode(form).encode()}
            scope = {'type': 'http', 'method': 'POST', 'scheme': 'https', 'server': ('calls.example.com', 443), 'path': path, 'query_string': b'', 'headers': []}
            with self.assertRaises(HTTPException) as error:
                await twilio_call.signed_form(Request(scope, receive), chats.secret())
            self.assertEqual(error.exception.status_code, 403)
            data = 'https://calls.example.com'+path+''.join(k+form[k] for k in sorted(form))
            sig = base64.b64encode(hmac.new(b'test-secret', data.encode(), hashlib.sha1).digest())
            scope['headers'] = [(b'x-twilio-signature', sig)]
            self.assertEqual(await twilio_call.signed_form(Request(scope, receive), chats.secret()), form)

    async def test_pstn_status_callback_does_not_require_optional_whatsapp_sender(self):
        from starlette.requests import Request
        from urllib.parse import urlencode
        cfg = {**self.cfg, 'whatsapp_from': ''}
        db.run("UPDATE integrations SET config=? WHERE kind='twilio'", (json.dumps(cfg),))
        path = f'/api/hooks/chat/{self.ws}/{chats.secret()}/twilio/voice/status/1'
        form = {'AccountSid': self.cfg['account_sid'], 'CallSid': 'CA'+'c'*32, 'CallStatus': 'completed'}
        data = 'https://calls.example.com' + path + ''.join(k+form[k] for k in sorted(form))
        sig = base64.b64encode(hmac.new(b'test-secret', data.encode(), hashlib.sha1).digest())
        async def receive():
            return {'type': 'http.request', 'body': urlencode(form).encode()}
        scope = {'type': 'http', 'method': 'POST', 'scheme': 'https', 'server': ('calls.example.com', 443),
                 'path': path, 'query_string': b'', 'headers': [(b'x-twilio-signature', sig)]}
        with db.use(self.ws):
            self.assertEqual(await twilio_call.signed_form(Request(scope, receive), chats.secret(), voice=True), form)

    async def test_status_callback_reconciles_call_without_local_media_session(self):
        with db.use(self.ws):
            cid = db.run("INSERT INTO calls(contact_id,status,mode,direction,ext_id,created_at) VALUES(?,?,?,?,?,?)",
                         (self.chat['contact_id'], 'ringing', 'pstn', 'out', 'CA'+'f'*32, db.now()))
            with mock.patch.object(twilio_call, 'signed_form', new_callable=mock.AsyncMock,
                                   return_value={'CallSid': 'CA'+'f'*32, 'CallStatus': 'completed', 'CallDuration': '1'}):
                response = await twilio_call.call_status(self.ws, 'secret', cid, None)
            row = db.q1('SELECT status,summary,duration,outcome,report FROM calls WHERE id=?', (cid,))
            self.assertEqual(response.status_code, 204)
            self.assertEqual(row['status'], 'completed')
            self.assertIn('no customer speech was captured', row['summary'])
            self.assertEqual(row['duration'], 1)
            self.assertEqual(row['outcome'], 'unresolved')
            self.assertEqual(json.loads(row['report'])['transcript_status'], 'unavailable')

    async def test_completed_callback_recovers_transcript_after_local_session_cleanup(self):
        from app import callengine
        with db.use(self.ws):
            cid = db.run("INSERT INTO calls(contact_id,status,mode,direction,ext_id,created_at,transcript) VALUES(?,?,?,?,?,?,?)",
                         (self.chat['contact_id'], 'failed', 'pstn', 'out', 'CA'+'e'*32, db.now(),
                          json.dumps([{'role': 'agent', 'text': 'Hello'}, {'role': 'customer', 'text': 'Call me tomorrow'}])))
            with mock.patch.object(twilio_call, 'signed_form', new_callable=mock.AsyncMock,
                                   return_value={'CallSid': 'CA'+'e'*32, 'CallStatus': 'completed', 'CallDuration': '12'}), \
                 mock.patch.object(callengine, 'finalize', new_callable=mock.AsyncMock) as finalize:
                response = await twilio_call.call_status(self.ws, 'secret', cid, None)
            row = db.q1('SELECT status,duration FROM calls WHERE id=?', (cid,))
            self.assertEqual(response.status_code, 204)
            self.assertEqual(row['status'], 'in_progress')
            self.assertEqual(row['duration'], 12)
            finalize.assert_awaited_once_with(cid, 'completed')

    async def test_foreign_workspace_cannot_attach_to_call(self):
        socket = SimpleNamespace(close=mock.AsyncMock(), accept=mock.AsyncMock())
        await twilio_call.media(socket, '../other-company', 'secret', 1)
        socket.close.assert_awaited_once_with(code=1008)
        socket.accept.assert_not_awaited()

    async def test_inbound_disabled_returns_reject(self):
        with db.use(self.ws), mock.patch.object(twilio_call, 'signed_form', new_callable=mock.AsyncMock,
                return_value={'From': 'whatsapp:+447700900123'}):
            response = await twilio_call.inbound(self.ws, 'secret', None)
            self.assertIn(b'<Reject', response.body)

    async def test_new_chat_prefers_twilio_and_sends_content_template(self):
        with mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock, return_value={'sid': 'SM-new'}) as send:
            status, result = await self.owner.post('/api/chats/new', {'phone': '+447700900999', 'template': 'HX'+'d'*32, 'params': ['Sam']})
            self.assertEqual(status, 200, result)
            with db.use(self.ws):
                self.assertEqual(chats._chat(result['chat_id'])['provider'], 'twilio')
            body = send.call_args.args[1]
            self.assertEqual(body['ContentSid'], 'HX'+'d'*32)
            self.assertEqual(json.loads(body['ContentVariables']), {'1': 'Sam'})

    async def test_permission_send_rechecks_closed_window(self):
        with db.use(self.ws), mock.patch.object(twilio_call, 'request', new_callable=mock.AsyncMock) as send:
            db.update('chats', self.chat['id'], last_in_at=db.now()-86401)
            with self.assertRaisesRegex(wa_call.CallError, 'active WhatsApp conversation'):
                await twilio_call.send_permission(self.chat)
            send.assert_not_awaited()
